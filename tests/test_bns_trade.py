"""Exports by commodity group from the BNS trade workbook — parsed from a workbook built
in the shape of the live file (regional blocks of 6-digit HS lines, [tonnes, unit, USD]
per month, national total in row 4); no network."""
import io
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from lib import validation  # noqa: E402
from fetchers import bns_trade  # noqa: E402

DS = {"id": "EXPORTS_VALUE_BY_COMMODITY_GROUP", "measure": "usd", "dictionary": "hs_export_groups"}
GROUPS = [{"code": "TOTAL", "name_ru": "всего", "prefixes": []},
          {"code": "WHEAT", "name_ru": "пшеница", "prefixes": ["1001"]},
          {"code": "FLAT_ROLLED", "name_ru": "прокат", "prefixes": ["7208", "7209"]},
          {"code": "CRUDE_OIL", "name_ru": "нефть", "prefixes": ["2709"]}]


MONTHS = ["январь", "февраль", "март", "апрель", "май", "июнь"]
# per line: (tonnes, usd) in January; month m adds m-1 to each
LINES = [("Акмолинская", None), ("100199", (100.0, 30.0)), ("720810", (5.0, 4.0)),
         ("Атырауская", None), ("270900", (1000.0, 500.0)), ("720890", (1.0, 1.0)), ("720916", (2.0, 2.0))]


def _workbook(national_override: dict | None = None) -> bytes:
    """Two regional blocks, six months. Row 4 = the national total of every line;
    `national_override` = {column index: value} to plant a wrong published total."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2025"
    ws.append(["Экспорт"])
    header, units = ["Код ТНВЭД", "Наименование", "Доп. ед."], ["тонн", "доп.", "тыс.долларов США"]
    for m in MONTHS:
        header += [f"{m} 2025 года", None, None]
        units += ["тонн", "доп.", "тыс.долларов США"]
    ws.append(header)
    ws.append(units)
    def row_vals(base):
        out = []
        for k in range(len(MONTHS)):
            out += [base[0] + k, 0, base[1] + k]
        return out
    tot = [0.0] * (3 * len(MONTHS))
    for _, base in LINES:
        if base:
            tot = [a + b for a, b in zip(tot, row_vals(base))]
    for col, v in (national_override or {}).items():
        tot[col - 3] = v
    ws.append(["Республики Казахстан", None, None] + tot)
    for code, base in LINES:
        ws.append([code, "name", None] + (row_vals(base) if base else []))
    wb.create_sheet("Метаданные").append(["x"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_groups_are_sums_of_their_hs_lines_and_total_is_the_national_row():
    out = bns_trade.parse_groups(_workbook(), GROUPS, DS)
    assert sorted(out["usd"]) == [f"2025-{m:02d}-01" for m in range(1, 7)]
    assert out["usd"]["2025-01-01"] == {"WHEAT": 30.0, "FLAT_ROLLED": 7.0, "CRUDE_OIL": 500.0, "TOTAL": 537.0}
    assert out["tonnes"]["2025-02-01"] == {"WHEAT": 101.0, "FLAT_ROLLED": 11.0, "CRUDE_OIL": 1001.0, "TOTAL": 1113.0}


def test_a_month_whose_regions_do_not_sum_to_the_national_total_is_skipped():
    content = _workbook(national_override={3: 1106.0})        # January tonnes total wrong by 0.2 % (true 1108)
    out = bns_trade.parse_groups(content, GROUPS, DS)
    assert "2025-01-01" not in out["tonnes"] and "2025-01-01" in out["usd"] and "2025-02-01" in out["tonnes"]


def test_a_non_six_digit_code_is_a_structural_change():
    wb = openpyxl.load_workbook(io.BytesIO(_workbook()))
    wb["2025"].append(["2709", "chapter subtotal", None] + [1.0, 0, 1.0] * len(MONTHS))
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(validation.StructuralChangeError) as exc:
        bns_trade.parse_groups(buf.getvalue(), GROUPS, DS)
    assert "4-digit" in str(exc.value)


def test_dictionary_prefixes_cover_the_model_groups_without_overlap():
    groups = bns_trade.load_groups("hs_export_groups")
    codes = [g["code"] for g in groups]
    assert len(codes) == len(set(codes)) and "TOTAL" in codes
    prefixes = [(p, g["code"]) for g in groups for p in g["prefixes"]]
    for p, code in prefixes:                      # no prefix is a prefix of another group's prefix
        for q, other in prefixes:
            if code != other:
                assert not q.startswith(p), f"{code} {p} overlaps {other} {q}"
    assert next(g for g in groups if g["code"] == "FLAT_ROLLED")["prefixes"] == ["7208", "7209", "7210", "7211", "7212"]
