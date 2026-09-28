"""RB_QUASIFISCAL_YTD (fetchers/minfin_quasifiscal.py): quasi-fiscal loans and equity through the
republican budget from the Statistical Bulletin «табл 8 (расх)». Parser tests on built fixtures
laid out as the bulletin (A function, B subfunction, C administrator, D programme, E Kazakh name,
F value, G Russian name); data tests on the processed file when it exists."""
import csv
import io
import sys
from pathlib import Path

import openpyxl
import pytest

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT / "scripts"))
from fetchers import minfin, minfin_quasifiscal as qf  # noqa: E402
from lib import validation  # noqa: E402

PROCESSED = REPO_ROOT / "data" / "processed" / "dims" / "rb_quasifiscal_ytd.csv"

BAITEREK_DBK = ("Кредитование АО «Национальный управляющий холдинг «Байтерек» с последующим кредитованием "
                "АО «Банк Развития Казахстана» для финансирования крупных проектов обрабатывающей промышленности")
IDF_VIA_DBK = ("Кредитование АО «Национальный управляющий холдинг «Байтерек» с последующим кредитованием "
               "АО «Фонд развития промышленности» через АО «Банк Развития Казахстана» по реализации в лизинг автобусов")
ACC = "Кредитование АО «Аграрная кредитная корпорация» для проведения мероприятий по поддержке субъектов АПК"
LOCAL = "Кредитование областных бюджетов, бюджетов городов республиканского значения, столицы на содействие развитию"
SAMRUK_EQ = "Увеличение уставного капитала АО «Фонд национального благосостояния «Самрук-Казына» для обеспечения"
INTL = "Приобретение акций международных финансовых организаций"
OTHER_EQ = "Увеличение уставного капитала АО «Казахстанский центр государственно-частного партнерства»"
SPACE = "Увеличение уставного капитала АО «Национальная компания «Қазақстан Ғарыш Сапары» для космического ракетного комплекса «Байтерек»"


def _row(func=None, abp=None, prg=None, value=None, ru=""):
    return [func, None, abp, prg, "қаз", value, ru]


def table8(period="январь-июль отчет 2026г.", loans=((229, 36, 50000.0, BAITEREK_DBK), (212, 262, 70000.0, ACC),
                                                        (None, 1, 30000.0, LOCAL)),
           assets=((243, 49, 20000.0, SAMRUK_EQ), (217, 6, 150.0, INTL)), loan_total=None, asset_total=None):
    rows = [[None, None, None, None, "Республикалық бюджет", None, "Республиканский бюджет"],
            [None, None, None, None, None, period, None],
            [None, None, None, None, "ШЫҒЫНДАР", 900000.0, "ЗАТРАТЫ"],
            _row(1, 217, None, 5000.0, "Министерство финансов"),
            _row(None, None, 1, 5000.0, "Услуги по исполнению бюджета"),
            [None, None, None, None, "БЮДЖЕТТІК КРЕДИТТЕР", loan_total if loan_total is not None else sum(v for *_, v, _ in loans),
             "БЮДЖЕТНЫЕ КРЕДИТЫ"]]
    for abp, prg, v, label in loans:
        if abp is not None:
            rows.append(_row(10, abp, None, v, "Администратор"))
        rows.append(_row(None, None, prg, v, label))
    rows.append([None, None, None, None, "ҚАРЖЫ АКТИВТЕРІН САТЫП АЛУ",
                 asset_total if asset_total is not None else sum(v for *_, v, _ in assets), "ПРИОБРЕТЕНИЕ ФИНАНСОВЫХ АКТИВОВ"])
    for abp, prg, v, label in assets:
        rows.append(_row(1, abp, prg, v, label))
    rows.append([None, None, None, None, "ҚАРЫЗДАРДЫ ӨТЕУ", 1.0, "ПОГАШЕНИЕ ЗАЙМОВ"])
    return rows


def table7(period="январь-июль отчет 2026г.", loans=150000.0, assets=20150.0):
    return [[None, "Показатели", "2025 г. отчет", period],
            [None, "x", 1.0, 2.0],
            [None, "Бюджетные кредиты", 999.0, loans],
            [None, "Приобретение финансовых активов", 999.0, assets]]


def table10(period="январь-июль отчет 2026г.", spec=50000.0, other=70000.0):
    return [[None, None, period, None],
            [None, 513, spec, "Бюджетные кредиты специализированным организациям"],
            [None, 519, other, "Прочие внутренние бюджетные кредиты"]]


def sheets(**kw):
    return {"табл 7": table7(**kw.pop("t7", {})), "табл 8 (расх)": table8(**kw.pop("t8", {})),
            "табл 10": table10(**kw.pop("t10", {}))}


def values(records):
    return {r["item_code"]: r["value"] for r in records}


# ---------------------------------------------------------------- recipients

@pytest.mark.parametrize("label, direct, final, family", [
    (BAITEREK_DBK, "BAITEREK", "DBK", "DEV_FINANCE"),
    (IDF_VIA_DBK, "BAITEREK", "IDF", "DEV_FINANCE"),                  # «через» DBK: a conduit
    ("Кредитование АО «НУХ «Байтерек» с последующим кредитованием АО «БРК-Лизинг»", "BAITEREK", "IDF", "DEV_FINANCE"),
    (ACC, "ACC", "ACC", "AGRI"),
    ("Кредитование АО «Аграрная кредитная корпорация» с последующим кредитованием АО «КазАгроФинанс»", "ACC", "KAF", "AGRI"),
    ("103 Увеличение уставного капитала АО «Фонд проблемных кредитов»", "FPK", "FPK", "FPK"),
    (SAMRUK_EQ, "SAMRUK", "SAMRUK", "SAMRUK"),
    ("Кредитование АО «ФНБ «Самрук-Қазына» с последующим кредитованием АО «НК «Қазақстан темір жолы»", "SAMRUK", "KTZ", "SAMRUK"),
    ("Кредитование АО «Жилищный строительный сберегательный банк «Отбасы банк» для займов", "OTBASY", "OTBASY", "HOUSING"),
])
def test_recipient_is_the_last_entity_before_via(label, direct, final, family):
    c = qf.classify(label)
    assert (c["class"], c["direct"], c["final"], c["family"]) == ("DI", direct, final, family)


@pytest.mark.parametrize("label, cls", [(LOCAL, "LOCAL_GOV"), ("Выполнение обязательств по государственным гарантиям", "GUARANTEE_CALL"),
                                        (INTL, "INTL_ORG"), (OTHER_EQ, "OTHER_QUASI"), (SPACE, "OTHER_QUASI"),
                                        ("Создание технопарка «Парк ядерных технологий» в городе Курчатове", "OTHER")])
def test_non_institution_classes(label, cls):
    assert qf.classify(label)["class"] == cls          # the space complex «Байтерек» is not the holding


def test_truncated_label_is_flagged_unless_known():
    label = ("Увеличение уставного капитала акционерного общества «Национальный управляющий холдинг «Байтерек» с "
             "последующим увеличением уставного капитала акционерного общества «Банк Развития Казахстана» с последующим "
             "увеличением уставного капитала акци")
    assert len(label) >= qf.TRUNCATED_AT
    assert qf.classify(label, "EQUITY", 229, 7, 2027)["truncated"] is True
    known = qf.classify(label, "EQUITY", 229, 7, 2026)
    assert (known["final"], known["truncated"]) == ("IDF", False)


# ---------------------------------------------------------------- one edition

def test_edition_items_add_up():
    d, records, warnings, programmes = qf.edition_records(sheets())
    v = values(records)
    assert d == "2026-07-01" and warnings == []
    assert v["LOAN"] == 150000.0 and v["LOAN.DI"] == 120000.0 and v["LOAN.LOCAL_GOV"] == 30000.0
    assert v["LOAN.DBK"] == 50000.0 and v["LOAN.ACC"] == 70000.0 and v["LOAN.VIA.BAITEREK"] == 50000.0
    assert v["LOAN.FAM.DEV_FINANCE"] == 50000.0 and v["LOAN.FAM.AGRI"] == 70000.0
    assert v["EQUITY"] == 20150.0 and v["EQUITY.SAMRUK"] == 20000.0 and v["EQUITY.INTL_ORG"] == 150.0
    assert set(v) == set(qf.ITEM_NAMES) and qf.identity_errors(v) == []
    assert [(p["abp"], p["prg"]) for p in programmes] == [(229, 36), (212, 262), (212, 1), (243, 49), (217, 6)]
    assert all(r["item_name"] == qf.ITEM_NAMES[r["item_code"]] for r in records)


def test_programmes_not_adding_up_stop_the_run():
    with pytest.raises(validation.StructuralChangeError, match="LOAN: programmes sum to 150000"):
        qf.edition_records(sheets(t8={"loan_total": 151000.0}))


def test_stale_table8_header_takes_the_edition_period():
    """2018-10 and 2020-07: «табл 8 (расх)» still headed with the previous month."""
    d, records, warnings, _ = qf.edition_records(sheets(t8={"period": "январь-июнь отчет 2026г."}))
    assert d == "2026-07-01" and values(records)["LOAN"] == 150000.0
    assert any("stale header" in w for w in warnings)


def test_header_without_otchet_and_string_values():
    """2018-03 / 2019-10 headers carry no «отчет»; 2017 values may be strings with spaces."""
    rows = table8(period="январь-март 2018 г.", loans=((229, 36, 50000.0, BAITEREK_DBK),), assets=())
    assert rows[7][6] == BAITEREK_DBK
    rows[7][5] = "  50\xa0000"
    parsed = qf.parse_table8(rows)
    assert parsed["period"] == (2018, 3) and parsed["items"][0]["value"] == 50000.0
    assert qf.section_errors(parsed) == []


def test_second_block_is_ignored():
    """The January 2014 edition appends a Jan-Dec 2013 block after its own."""
    rows = table8() + table8(period="январь-декабрь отчет 2013г.", loans=((212, 23, 9e9, ACC),))
    parsed = qf.parse_table8(rows)
    assert parsed["sections"]["LOAN"] == 150000.0 and len(parsed["items"]) == 5


def test_january_edition_without_sections_is_all_zero():
    rows = table8(period="январь отчет 2026г.")[:5] + [[None, None, None, None, "ҚАРЫЗДАРДЫ ӨТЕУ", 1.0, "ПОГАШЕНИЕ ЗАЙМОВ"]]
    d, records, _, programmes = qf.edition_records({"табл 8 (расх)": rows})
    assert d == "2026-01-01" and programmes == [] and all(r["value"] == 0.0 for r in records)


def test_table7_and_table10_mismatches_warn():
    _, _, warnings, _ = qf.edition_records(sheets(t7={"loans": 140000.0}, t10={"other": 60000.0}))
    assert any("LOAN: табл 8 150000.000 vs табл 7 140000.000" in w for w in warnings)
    assert any("LOAN.DI: 120000.000 vs табл 10 513+519 110000.000" in w for w in warnings)


def test_known_table7_error_is_tagged():
    kw = {"period": "январь-февраль отчет 2021г."}
    _, _, warnings, _ = qf.edition_records(sheets(t7={**kw, "loans": 80337.874}, t8=kw, t10=kw))
    assert any("known табл 7 error" in w for w in warnings)


def test_december_edition_reads_the_report_column_of_table7():
    kw = {"period": "январь-декабрь отчет 2025г."}
    t7 = [[None, "Показатели", "2024 г. отчет", "2025 г. отчет"], [None, "Бюджетные кредиты", 1.0, 150000.0],
          [None, "Приобретение финансовых активов", 1.0, 20150.0]]
    assert qf.table7_totals(t7, 2025, 12) == {"LOAN": 150000.0, "EQUITY": 20150.0}
    s = sheets(t8=kw, t10=kw)
    s["табл 7"] = t7
    assert qf.edition_records(s)[2] == []


def test_unknown_loan_recipient_warns():
    loans = ((229, 36, 50000.0, BAITEREK_DBK), (229, 90, 10.0, "Кредитование АО «Новый оператор» для чего-то"))
    _, _, warnings, _ = qf.edition_records(sheets(t7={"loans": 50010.0}, t8={"loans": loans}, t10={"other": 0.0}))
    assert any("loan to an unknown recipient (OTHER_QUASI)" in w for w in warnings)


def test_ytd_decrease_and_zero_series():
    recs = [{"date": d, "item_code": "LOAN.DBK", "value": v} for d, v in
            (("2025-11-01", 10.0), ("2025-12-01", 8.0), ("2026-01-01", 0.0))]
    assert qf.ytd_warnings(recs) == ["LOAN.DBK: year to date falls from 10.0 (2025-11) to 8.0 (2025-12)"]
    recs.append({"date": "2026-01-01", "item_code": "LOAN.FPK", "value": 0.0})
    assert {r["item_code"] for r in qf.drop_zero_series(recs)} == {"LOAN.DBK"}


# ---------------------------------------------------------------- fetch, as the backfill cache drives it

class _Book:
    def __init__(self, s):
        self.sheets, self.sheetnames = s, list(s)


def test_fetch_newest_edition_wins_and_keeps_history(monkeypatch):
    books = {"/a.xlsx": sheets(),                                                   # Jan-Jul 2026
             "/b.xlsx": sheets(t7={"loans": 1.0}),                                  # the same period, older document
             "/c.xlsx": {"табл 1": [[1]]}}                                          # no table 8
    docs = [{"id": i, "created_date": f"2026-09-0{i}", "full_text": [{"document": p}]} for i, p in ((3, "/a.xlsx"), (2, "/b.xlsx"), (1, "/c.xlsx"))]
    monkeypatch.setattr(minfin, "_bulletin_documents", lambda: docs)
    monkeypatch.setattr(minfin, "_list_documents", lambda **kw: [])
    monkeypatch.setattr(minfin, "_download", lambda p: p.encode())
    monkeypatch.setattr(minfin, "_open_workbook", lambda c, hint: ("sheets", _Book(books[c.decode()])))
    monkeypatch.setattr(minfin, "_iter_rows", lambda kind, wb, name=None: iter(wb.sheets[name]))
    stored = [{"date": "2026-06-01", "region": "national", "item_code": "LOAN", "item_name": "x", "value": "100000.0"},
              {"date": "2026-07-01", "region": "national", "item_code": "LOAN", "item_name": "x", "value": "1.0"}]
    monkeypatch.setattr(qf.dims, "load_processed", lambda ds_id: stored)
    records, info = qf.fetch({"id": qf.DATASET_ID, "note": "n"})
    loan = {r["date"]: r["value"] for r in records if r["item_code"] == "LOAN"}
    assert loan == {"2026-06-01": 100000.0, "2026-07-01": 150000.0}
    assert "LOAN.FPK" not in {r["item_code"] for r in records}                     # zero everywhere: dropped
    assert any("no «табл 8 (расх)»" in w for w in info["warnings"])
    assert not any("табл 7" in w for w in info["warnings"])                         # the older duplicate is not read


def test_untitled_bulletin_is_added(monkeypatch):
    listed = [{"id": 964016, "title": "(Is developed in Kazakh and Russian, while English translation is not provided)",
               "created_date": "2026-02-04", "full_text": [{"document": "/uploads/x.xlsx"}]},
              {"id": 5, "title": "Other", "created_date": "2026-02-05", "full_text": [{"document": "/uploads/y.xlsx"}]}]
    monkeypatch.setattr(minfin, "_bulletin_documents", lambda: [])
    monkeypatch.setattr(minfin, "_list_documents", lambda **kw: listed)
    assert [d["id"] for d in qf._documents()] == [964016]


def test_read_workbook_from_xlsx(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in {**sheets(), "табл 9": [[1]]}.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    f = tmp_path / "b.xlsx"
    f.write_bytes(buf.getvalue())
    s = qf._read_workbook(f)
    assert set(s) == {"табл 7", "табл 8 (расх)", "табл 10"}
    assert values(qf.edition_records(s)[1])["LOAN.DI"] == 120000.0


# ---------------------------------------------------------------- processed data (after the integrator's backfill)

@pytest.mark.skipif(not PROCESSED.exists(), reason="RB_QUASIFISCAL_YTD not built yet")
def test_processed_december_totals():
    with PROCESSED.open(encoding="utf-8") as f:
        v = {(r["date"], r["item_code"]): float(r["value"]) for r in csv.DictReader(f)}
    di = [60.0, 58.9, 68.9, 182.5, 194.1, 146.5, 246.5, 245.2, 284.7, 577.3, 410.3, 274.0]
    for year, bn in zip(range(2013, 2025), di):
        assert v[(f"{year}-12-01", "LOAN.DI")] / 1000 == pytest.approx(bn, abs=0.05)
    assert v[("2014-12-01", "EQUITY.FPK")] == pytest.approx(250000.0)
    assert v[("2020-12-01", "LOAN")] == pytest.approx(338215.473, abs=0.001)
    for (d, item), x in v.items():
        if item in ("LOAN", "EQUITY"):
            parts = sum(v.get((d, f"{item}.{c}"), 0.0) for c in ["DI"] + qf.REMAINDERS)
            assert parts == pytest.approx(x, abs=0.5), (d, item)
