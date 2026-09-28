"""Balance sheets of the state development institutions and quasi-fiscal holdings from their
statements on KASE (added 2026-09-27 for the quasi-fiscal block; extended the same day to the
Problem Loans Fund, KazAgro and Samruk-Kazyna).

Issuers (config/dims.yaml `issuers`; the key is the KASE code unless ORG_CODE maps it):
  BTRK -- «Байтерек» holding (BIN 130540020197, created 2013), consolidated: DBK, Damu, Otbasy
          bank, KazakhExport, Kazakhstan Housing Company, IDF, QIC and, from 2021 (KazAgro merged
          in), the Agrarian Credit Corporation and KazAgroFinance. KASE xlsx from 2019.
  BRKZ -- Development Bank of Kazakhstan, consolidated (a BTRK subsidiary).
  KFUS -- АО «Казахстанский фонд устойчивости» (BIN 170940012405), created by NBK resolution №130
          of 30.06.2017 to run the bank-recapitalisation programme; absorbed ИО «Баспана» on
          25.12.2019 (total assets 1 536.6 bn at 2019-12-31 -> 2 909.1 bn at 2020-03-31, a break).
          It is NOT the Problem Loans Fund and not its successor.
  AGKK -- Agrarian Credit Corporation (a KazAgro, then BTRK, subsidiary).
  FPKR -- АО «Фонд проблемных кредитов» (Problem Loans Fund, ФПК; BIN 120140005984, registered
          11.01.2012 under NBK resolution №53 of 30.05.2011, 100% state). Its bonds (FPKRb1 450 bn,
          b2 604 bn, b3 17.5 bn) were annulled on 21.09.2021 and it stopped filing on KASE. The
          issuer page and the documents API now list nothing, but 11 consolidated xlsx remain on
          the server (9M-2018 .. 6M-2021, 2017-12-31 as a comparative): STATIC_DOCS. Negative
          equity since 2018 (-116.4 bn at 2018-12-31, -716.3 bn at 2019-12-31); the equity line
          «Резерв по условному распределению» holds the cumulative transfers to banks.
  KZAG -- НУХ «КазАгро» (BIN 070140002180), consolidated (AKK, KazAgroFinance, ФФПСХ ...);
          delisted in May 2021 when it was merged into BTRK. 23 consolidated quarterly xlsx
          remain on the server (9M-2013 .. 9M-2020, 2012-12-31 as a comparative): STATIC_DOCS.
          Total assets 548.1 bn at 2012-12-31, 1 474.3 bn at 2020-09-30.
  SKKZ -- ФНБ «Самрук-Казына» (BIN 081140000436), consolidated (dominated by KMG, KTZ, KAP). xlsx
          only from 9M-2022 (2021-12-31 comparative), in millions; the H1/annual statements of
          2011-2025 are scanned PDFs.
  SKKZ_SEP -- Samruk-Kazyna's own (separate, «не консолидированный») balance sheet: the Minfin
          Form 1 «неполная финансовая отчетность» files *_nb_rus.xlsx, 2013-12 .. 2016-03
          (2012-12-31 as a comparative). Loans to and investments in subsidiaries, own borrowings.
Never add a subsidiary to its parent: BTRK consolidates BRKZ and AGKK (from 2021), KZAG
consolidates AGKK (to 2020). The quasi-fiscal aggregate is BTRK + KZAG (to 2020Q4) + FPKR + KFUS
(+ SKKZ_SEP).

Discovery: https://kase.kz/api/companies/documents/?language=ru&org_code=<CODE> (JSON; the
trailing slash and `language` are required) lists every document an issuer filed. Statements are
the xlsx/xls whose name contains «отчетност» («Финансовая отчетность за январь–<месяц> YYYY
года», «Годовая … финансовая отчетность за YYYY год», «Неполная финансовая отчетность …» for the
Minfin-form *_nb_* files). Delisted issuers return [], so FPKR and KZAG fall back to STATIC_DOCS
(file names found by probing the name pattern; every other candidate 404s). .xlsb (one Damu
file) is skipped: pyxlsb is not installed. File names are not fully regular
(skkzfm3_2025_rus.xlsx is consolidated although it lacks «cons»; kfusfm1_2021_rus_2.xlsx).

KASE fin-data API https://kase.kz/api/companies/fin-data/<CODE>/: total assets, liabilities and
equity per reporting date (`change_date` is the day AFTER the balance date; units thnd|mln).
Used (a) as a unit cross-check of every statement date it covers -- a factor of 100 or more raises
StructuralChangeError (a gap above 1% is a warning: KFUS restated 2019-12-31 to 1 477.3 bn in 9M-2020) -- and (b) for the issuers in `fin_data` as separate items FD_TOTAL_ASSETS,
FD_TOTAL_LIABILITIES, FD_EQUITY: BTRK year ends 2015-2017 and quarterly from 2018-06 (the xlsx
start in 2019), SKKZ half-yearly from 2007 (liabilities null before 2016). FPKR and KZAG return
null. Verified: BTRK 2015-12-31 assets 3 460.3 bn, 2017-12-31 4 432.6 bn; SKKZ 2026-06-30
46 665 321 mln = the xlsx.

Parsing (parse_balance_sheet): .xls is read with xlrd, .xlsx with openpyxl. The balance-sheet
sheet is the first one with an «Итого/Всего активов» row (or the Minfin «Баланс (строка 100 +
строка 101 + строка 200)» row, or failing both, «Итого обязательства и капитал», which equals
assets). Units: the first cell naming them («тыс./тысяч» -> 1e3, «млн/миллион» -> 1e6); if none
does, thousands are assumed and noted (FPKR 2019-2021 and KZAG 2015-2017 have no unit line;
thousands confirmed by the totals' continuity with the files that state it). Date columns are the
header cells that read as dates («31 марта 2026 г.», «31.12.2019», «30 сентября» over «2014
года»); the Minfin form's «На конец / На начало отчетного периода» take the reporting date from
the title («за период с 01.01.2015 по 30.09.2015», «по состоянию на "01" апреля 2019 года») and
Dec 31 of the previous year. A balance date on the 1st of a month is the previous day's close
(«на 01.04.2021» -> 2021-03-31). Titles and headers carry typos (KFUS 9M-2022 heads its September
column «30 июня 2022»; KZAG 9M-2014 is titled «На 30 июня 2014 года»), so the reporting period in
the file name (fm1..fm3 = Mar/Jun/Sep, fm4 or f = Dec) decides between them; a comparative column
that is not a year end is dropped. Labels lose «(сумма строк …)»/«(строка …)» suffixes. Lines
(LINES) are matched on the label: totals take the first matching row; other lines SUM every
matching row in the balance sheet (current + long-term «Займы», tenge bonds + Eurobonds). Total
liabilities absent -> short-term + long-term totals (Minfin form). A column is kept only if
TOTAL_ASSETS = TOTAL_LIABILITIES + EQUITY within 0.5%; where two columns carry the same date
(SK Form 1 has an empty «На конец» column), the first one passing the identity is used.
"""
from __future__ import annotations

import io
import json
import re
import sys
import time
from datetime import date, datetime, timedelta
from pathlib import Path

import requests

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from fetchers import imf_dims  # noqa: E402
from lib import dims, raw_store, validation  # noqa: E402

DOCS_URL = "https://kase.kz/api/companies/documents/?language=ru&org_code={code}"
FIN_DATA_URL = "https://kase.kz/api/companies/fin-data/{code}/"
FILE_BASE = "https://kase.kz"
HEADERS = {"User-Agent": "Mozilla/5.0 (compatible; KZEconDataPipeline/1.0; +https://github.com/)"}
MONTHS = {"январ": 1, "феврал": 2, "март": 3, "апрел": 4, "ма": 5, "июн": 6, "июл": 7, "август": 8,
          "сентябр": 9, "октябр": 10, "ноябр": 11, "декабр": 12}
# dims.yaml issuer key -> KASE org code, where they differ (the separate SK statements are a
# second series of the same issuer)
ORG_CODE = {"SKKZ_SEP": "SKKZ"}
# Issuers whose statements must say «консолид» (they also file separate statements)
CONSOLIDATED_ONLY = {"BTRK", "BRKZ"}


def _static(code: str, names: list[str]) -> list[tuple[str, str]]:
    return [(f"Финансовая отчетность (файл на сервере KASE, эмитент исключен): {n}", f"/files/emitters/{code}/{n}")
            for n in names]


# Delisted issuers: the documents API returns [], but these files are still served (checked
# 2026-09-27; kzagfm{1,2}_2013, kzagfm4_{2013,2015..}, kzagf_*, kzagfm1_2021, fpkrfm{1,2}_2018,
# fpkrf_2018 (.xlsx), fpkrfm3_2021 and fpkrf_2021 all 404).
STATIC_DOCS = {
    "FPKR": _static("FPKR", ["fpkrfm3_2018_cons_rus.xlsx"]
                    + [f"fpkr{p}_{y}_cons_rus.xlsx" for y in (2019, 2020) for p in ("fm1", "fm2", "fm3", "f")]
                    + ["fpkrfm1_2021_cons_rus.xlsx", "fpkrfm2_2021_cons_rus.xlsx"]),
    "KZAG": _static("KZAG", ["kzagfm3_2013_cons_rus.xlsx"]
                    + [f"kzagfm{q}_2014_cons_rus.xlsx" for q in (1, 2, 3, 4)]
                    + [f"kzagfm{q}_{y}_cons_rus.xlsx" for y in range(2015, 2021) for q in (1, 2, 3)]),
}

# (code, Russian name, regex on the normalised label). Totals (TOTALS) take the first matching
# row, other lines sum every matching row; a label is assigned to the first line that matches.
LINES = [
    ("TOTAL_ASSETS", "Итого активов", r"^(итого|всего):? актив"),
    ("TOTAL_LIABILITIES", "Итого обязательств", r"^(итого|всего):? обязательств[оа]?$"),
    ("EQUITY", "Итого капитала", r"^(итого|всего):? (собственн\w* )?капитал\w*$"),
    ("LOANS_CUSTOMERS", "Кредиты, выданные клиентам", r"^(кредиты|займы),? (выданные |предоставленные )?клиентам"),
    ("LOANS_BANKS", "Кредиты, выданные банкам и финансовым институтам",
     r"^(кредиты|займы),? (выданные |предоставленные )?(банкам|кредитным|финансовым)"),
    ("LOANS_ISSUED", "Займы выданные (и чистые инвестиции в финансовую аренду)", r"^займы выданные"),
    ("FIN_LEASE", "Дебиторская задолженность по финансовой аренде", r"^дебиторская задолженность по (договорам )?финансов\w* аренд"),
    ("ACQUIRED_CLAIMS", "Права требования (приобретенные)", r"^права требовани"),
    ("DEBT_SECURITIES", "Выпущенные долговые ценные бумаги",
     r"^выпущенные (долговые ценные бумаги|облигации|еврооблигации)|^долговые ценные бумаги выпущенные"),
    ("GOV_LOANS", "Займы и средства от Правительства РК",
     r"^(займы|кредиты|средства),? (полученные )?от правительства|^(займы|кредиты|средства) правительства"
     r"|^задолженность перед правительством"),
    ("PARENT_LOANS", "Займы и задолженность перед акционером (материнской компанией)",
     r"^задолженность перед (акционером|материнской)|^займы,? (полученные )?от (материнской|акционера)"),
    ("BORROWINGS", "Займы полученные и средства кредитных учреждений",
     r"^(займы|займы полученные|средства кредитных (учреждений|организаций))$"),
    ("GOV_SUBSIDIES", "Государственные субсидии (отложенный доход)", r"^государственные субсидии(?! к получению)"),
    ("CUSTOMER_FUNDS", "Средства клиентов", r"^(средства|текущие счета и вклады) клиентов"),
    ("DISTRIBUTION_RESERVE", "Резерв по условному распределению (капитал)", r"^резерв по условному распределению"),
    # helpers, not stored
    ("_ST_LIAB", "", r"^(итого|всего):? краткосрочных обязательств$"),
    ("_LT_LIAB", "", r"^(итого|всего):? долгосрочных обязательств$"),
    ("_TL_EQ", "", r"^(итого|всего):? (обязательств\w*|капитал\w*) и (капитал\w*|обязательств\w*)$"),
]
_LINES = [(c, n, re.compile(r)) for c, n, r in LINES]
TOTALS = {"TOTAL_ASSETS", "TOTAL_LIABILITIES", "EQUITY", "_ST_LIAB", "_LT_LIAB", "_TL_EQ"}
FD_LINES = [("FD_TOTAL_ASSETS", "Итого активов (ключевые показатели KASE)", "aggregate_assets"),
            ("FD_TOTAL_LIABILITIES", "Итого обязательств (ключевые показатели KASE)", "total_liabilities"),
            ("FD_EQUITY", "Итого капитала (ключевые показатели KASE)", "own_capital")]
OLE_MAGIC = b"\xd0\xcf\x11\xe0"


def _norm(text) -> str:
    return re.sub(r"\s+", " ", str(text or "").replace("ё", "е")).strip().lower()


def _label(text) -> str:
    """Normalised row label without the Minfin form's «(сумма строк …)» / «(строка …)» suffixes."""
    s = _norm(text)
    if re.match(r"^баланс \(строк[аи]? ?100", s):
        return "итого активов"
    if re.match(r"^баланс \(строк[аи]? ?300", s):
        return "итого обязательства и капитал"
    s = re.sub(r"\s*\((сумма строк|строк[аи]?|стр\.)[^)]*\)", " ", s)
    return re.sub(r"\s+", " ", s).strip(" :")


def _day(y: int, m: int, d: int) -> str | None:
    try:
        return date(y, m, d).isoformat()
    except ValueError:  # «31.06.2022» in a malformed header
        return None


def _as_date(cell) -> str | None:
    if isinstance(cell, datetime):
        return cell.date().isoformat()
    s = _norm(cell)
    m = re.search(r"(\d{1,2})\.(\d{2})\.(\d{4})", s)
    if m:
        return _day(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    m = re.search(r"(\d{1,2})\s+([а-я]+)\s+(\d{4})", s)
    if m:
        month = next((v for k, v in MONTHS.items() if m.group(2).startswith(k)), None)
        if month:
            return _day(int(m.group(3)), month, int(m.group(1)))
    return None


def _close(d: str) -> str:
    """A balance «на 01.04.2021» is the close of 2021-03-31."""
    x = date.fromisoformat(d)
    return (x - timedelta(days=1)).isoformat() if x.day == 1 else d


def _title_date(rows: list[list]) -> str | None:
    """Reporting date named in the title rows: «за период с 01.01.2015 по 30.09.2015»,
    «по состоянию на "01" апреля 2019 года», «на 31 марта 2026 года»."""
    for k, r in enumerate(rows[:25]):
        for c in r:
            if not isinstance(c, str):
                continue
            s = _norm(c).replace('"', "").replace("«", "").replace("»", "")
            m = re.search(r"период с \d{1,2}\.\d{2}\.\d{4} по (\d{1,2})\.(\d{2})\.(\d{4})", s)
            if m:
                return _day(int(m.group(3)), int(m.group(2)), int(m.group(1)))
            if (k < 10 or "по состоянию на" in s) and re.search(r"(^|\s)(на|за)\s+\d", s) and _as_date(s):
                return _close(_as_date(s))
    return None


def _period_end(link: str) -> str | None:
    """Balance date of a statement from its file name: …fm3_2024… -> 2024-09-30; …f_2024… -> 2024-12-31."""
    y, q = _order(link)
    if not y:
        return None
    return (date(y + (q == 4), 1 if q == 4 else 3 * q + 1, 1) - timedelta(days=1)).isoformat()


def _number(cell) -> float | None:
    if isinstance(cell, bool):
        return None
    if isinstance(cell, (int, float)):
        return float(cell)
    s = str(cell or "").strip().replace("\xa0", "").replace(" ", "")
    if s in ("", "-", "–", "—"):
        return 0.0 if s else None
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):  # 24.297.440 (thousands separated by dots)
        s = s.replace(".", "")
    s = s.replace(",", ".")
    try:
        v = float(s)
    except ValueError:
        return None
    return -v if neg else v


def _sheets(content: bytes) -> list[tuple[str, list[list]]]:
    """(sheet name, rows) of an .xlsx (openpyxl) or .xls (xlrd) workbook; dates as datetime."""
    if content[:4] == OLE_MAGIC:
        import xlrd
        book = xlrd.open_workbook(file_contents=content)
        out = []
        for sh in book.sheets():
            rows = []
            for i in range(min(sh.nrows, 300)):
                row = []
                for j in range(sh.ncols):
                    cell = sh.cell(i, j)
                    v = None if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK) else cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            v = datetime(*xlrd.xldate_as_tuple(cell.value, book.datemode))
                        except (ValueError, xlrd.xldate.XLDateError):
                            pass
                    row.append(v)
                rows.append(row)
            out.append((sh.name, rows))
        return out
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True, max_row=300)]) for ws in wb.worksheets]


def _row_label(r: list) -> str:
    return next((_label(c) for c in r[:4] if isinstance(c, str) and re.search(r"[а-я]{3}", _norm(c))), "")


def _line(label: str) -> str | None:
    return next((code for code, _, rx in _LINES if rx.search(label)), None)


def _date_columns(rows: list[list], title: str | None) -> dict[int, str]:
    for k, r in enumerate(rows[:25]):
        found: dict[int, str] = {}
        for i, c in enumerate(r):
            s = _norm(c) if isinstance(c, str) else ""
            if title and re.match(r"^на конец\s+отчетного", s):
                found[i] = title
            elif title and re.match(r"^на начало\s+отчетного", s):
                found[i] = f"{int(title[:4]) - 1}-12-31"
            elif _as_date(c):
                found[i] = _as_date(c)
        if len(found) < 2 and k + 1 < len(rows):  # «На 31 марта» over «2019 года (неаудировано)»
            nxt = rows[k + 1]
            two = {i: _as_date(f"{c} {nxt[i] if i < len(nxt) else ''}") for i, c in enumerate(r) if isinstance(c, str)}
            two = {i: d for i, d in two.items() if d}
            if len(two) > len(found):
                found = two
        if len(found) >= 2:
            return {i: _close(d) for i, d in found.items()}
    # a statement with no comparative (FPKR 9M-2018): one date heading a column right of the labels
    for r in rows[:25]:
        found = {i: _as_date(c) for i, c in enumerate(r)
                 if i > 0 and _as_date(c) and not re.search(r"№|приказ|^от ", _norm(c))}
        if len(found) == 1:
            return {i: _close(d) for i, d in found.items()}
    return {}


def parse_balance_sheet(content: bytes, expected: str | None = None) -> tuple[dict[str, dict[str, float]], str]:
    """{date: {line code: bn KZT}} for the columns that pass the balance identity, and a note on
    what was assumed or skipped. `expected` is the balance date the file name gives (_period_end)."""
    sheets = _sheets(content)
    pick = None
    for want in ("TOTAL_ASSETS", "_TL_EQ"):
        pick = next((rows for _, rows in sheets if any(_line(_row_label(r)) == want for r in rows)), None)
        if pick is not None:
            break
    if pick is None:
        return {}, "no sheet with a total-assets row"
    rows = pick
    notes = []
    scale = None
    for r in rows[:20]:
        for c in r:
            s = _norm(c) if isinstance(c, str) else ""
            if re.search(r"тыс\.|тыс |тыс$|тысяч", s):
                scale = 1e3
            elif re.search(r"млн|миллион", s):
                scale = 1e6
            if scale:
                break
        if scale:
            break
    if scale is None:
        scale = 1e3
        notes.append("no unit line, thousands assumed")
    title = _title_date(rows)
    cols = _date_columns(rows, title or expected)  # «за первый квартал 2014 года»: the file name's date
    if not cols:
        return {}, "no header row with dates"
    # The current column is the latest date. Titles and headers both carry typos, so the file
    # name's period decides when it is given; otherwise the title does.
    current = max(cols, key=lambda i: cols[i])
    head = cols[current]
    if expected and expected in (head, title):
        new = expected
    elif expected and head[:4] != expected[:4]:
        return {}, f"header {head} / title {title} not in the file name's year ({expected})"
    elif expected:
        new = head
        notes.append(f"header {head} / title {title} differ from the file name's {expected}")
    else:
        new = title or head
    if new != head:
        notes.append(f"header {head} read as the {'title' if new == title else 'file name'}'s {new}")
        for i in [i for i in cols if cols[i] == head]:
            cols[i] = new
    current_date = new
    for i in [i for i in cols if cols[i] != current_date and not cols[i].endswith("-12-31")]:
        notes.append(f"comparative column headed {cols[i]} dropped (not a year end)")
        del cols[i]
    # rows of the balance sheet: up to the last of the first total-assets / liabilities / equity rows
    labels = [_line(_row_label(r)) for r in rows]
    firsts = [labels.index(c) for c in ("TOTAL_ASSETS", "TOTAL_LIABILITIES", "EQUITY", "_LT_LIAB", "_TL_EQ") if c in labels]
    end = max(firsts) + 1
    per_col: dict[int, dict[str, float]] = {i: {} for i in cols}
    for r, hit in zip(rows[:end], labels[:end]):
        if hit is None:
            continue
        for i in cols:
            v = _number(r[i]) if i < len(r) else None
            if v is None:
                continue
            if hit in TOTALS:
                per_col[i].setdefault(hit, v * scale / 1e9)
            else:
                per_col[i][hit] = per_col[i].get(hit, 0.0) + v * scale / 1e9
    for v in per_col.values():
        if "TOTAL_LIABILITIES" not in v and "_ST_LIAB" in v and "_LT_LIAB" in v:
            v["TOTAL_LIABILITIES"] = v["_ST_LIAB"] + v["_LT_LIAB"]  # Minfin form: no total row
        if not v.get("TOTAL_ASSETS") and v.get("_TL_EQ"):
            v["TOTAL_ASSETS"] = v["_TL_EQ"]
            if "total assets from" not in " ".join(notes):
                notes.append("total assets from «итого обязательства и капитал»")
    kept, skipped = {}, []
    for i in sorted(cols):
        d, v = cols[i], per_col[i]
        if d in kept:
            continue
        a, li, e = v.get("TOTAL_ASSETS"), v.get("TOTAL_LIABILITIES"), v.get("EQUITY")
        if a and li is not None and e is not None and abs(a - li - e) <= 0.005 * abs(a):
            kept[d] = {k: x for k, x in v.items() if not k.startswith("_")}
            if d in skipped:
                skipped.remove(d)
        elif d not in skipped:
            skipped.append(d)
    if skipped:
        notes.append(f"columns failing assets = liabilities + equity: {skipped}")
    return kept, "; ".join(notes)


def parse_fin_data(rows) -> dict[str, dict[str, float]]:
    """{balance date: {FD line: bn KZT}} from the KASE fin-data API rows."""
    if not isinstance(rows, list) or (rows and not all(isinstance(r, dict) and "change_date" in r for r in rows)):
        raise validation.StructuralChangeError(f"kase_ifrs: fin-data API returned {str(rows)[:200]}")
    out: dict[str, dict[str, float]] = {}
    for r in rows:
        scale = {"thnd": 1e3, "mln": 1e6}.get(r.get("units"))
        if scale is None:
            raise validation.StructuralChangeError(f"kase_ifrs: fin-data units {r.get('units')!r}")
        d = (date.fromisoformat(r["change_date"][:10]) - timedelta(days=1)).isoformat()
        vals = {code: r[key] * scale / 1e9 for code, _, key in FD_LINES if isinstance(r.get(key), (int, float))}
        if vals:
            out[d] = vals
    return out


def keep_document(issuer: str, name: str, link: str) -> bool:
    """A statement of this issuer: xlsx/xls, «отчетност» in the title; the Minfin-form *_nb_*
    files only for the *_SEP series and never for the others; consolidated only where required."""
    f, low = Path(link).name.lower(), name.lower()
    if not f.endswith((".xlsx", ".xls")) or "отчетност" not in low or "700" in f:
        return False
    if ("_nb_" in f) != issuer.endswith("_SEP"):
        return False
    return "консолид" in low or issuer not in CONSOLIDATED_ONLY


def documents(issuer: str) -> list[tuple[str, str]]:
    code = ORG_CODE.get(issuer, issuer)
    listed = json.loads(_get(DOCS_URL.format(code=code)))
    if not isinstance(listed, list):
        raise validation.StructuralChangeError(f"kase_ifrs: documents API for {code} returned {str(listed)[:200]}")
    docs = {(d["name"], d["link"]) for d in listed if keep_document(issuer, d.get("name") or "", d.get("link") or "")}
    if not docs:
        docs = set(STATIC_DOCS.get(issuer, []))
    if not docs:
        raise validation.StructuralChangeError(f"kase_ifrs: no financial statements listed for {code} at {DOCS_URL.format(code=code)}")
    return sorted(docs, key=lambda x: x[1])


def _get(url: str, missing_ok: bool = False) -> bytes | None:
    for attempt in range(4):
        try:
            r = requests.get(url, headers=HEADERS, timeout=120)
            if missing_ok and r.status_code == 404:
                return None
            r.raise_for_status()
            return r.content
        except requests.RequestException:  # kase.kz resets connections under load: retry serially
            if attempt == 3:
                raise
            time.sleep(2 ** (attempt + 1))
    raise AssertionError


def fetch(ds: dict) -> tuple[list[dict], dict]:
    fresh = imf_dims.stored_if_fresh(ds, "kase", ds["id"])
    stored = {r["item_code"].split(".")[0] for r in fresh[0]} if fresh else set()
    if fresh and stored >= set(ds["issuers"]) and all(f"{c}.FD_TOTAL_ASSETS" in {r["item_code"] for r in fresh[0]}
                                                      for c in ds.get("fin_data", [])):
        return fresh  # an issuer added to the config is downloaded at once, not after refresh_days
    today = date.today()
    names = {c: n for c, n, _ in LINES} | {c: n for c, n, _ in FD_LINES}
    best: dict[tuple[str, str], tuple[str, float]] = {}  # (issuer.line, date) -> (link, value) of the latest statement
    warnings, files = [], 0
    for issuer in ds["issuers"]:
        for _, link in documents(issuer):
            content = _get(FILE_BASE + link, missing_ok=True)
            if content is None:  # listed but not served (skkzfm3_2013_nb_rus.xlsx)
                warnings.append(f"{Path(link).name}: 404")
                continue
            files += 1
            ext = "xls" if content[:4] == OLE_MAGIC else "xlsx"
            raw_store.save_raw_bytes("kase", f"ifrs_{Path(link).stem}", today, ext, content)
            try:
                sheet, note = parse_balance_sheet(content, expected=_period_end(link))
            except Exception as exc:  # a corrupt file must not stop the other issuers
                sheet, note = {}, f"unreadable: {exc}"
            if note:
                warnings.append(f"{Path(link).name}: {note}")
            for d, lines in sheet.items():
                for line, v in lines.items():
                    key = (f"{issuer}.{line}", d)
                    if key not in best or (_order(link), link) > (_order(best[key][0]), best[key][0]):
                        best[key] = (link, v)
    # KASE key indicators: a unit check of the statements, and FD_* items for ds["fin_data"]
    for issuer in ds["issuers"]:
        if issuer in ORG_CODE or issuer in STATIC_DOCS:  # SEP is not the consolidated FD; delisted return null
            continue
        content = _get(FIN_DATA_URL.format(code=issuer))
        raw_store.save_raw_bytes("kase", f"findata_{issuer.lower()}", today, "json", content)
        fd = parse_fin_data(json.loads(content) or [])
        for d, v in fd.items():
            got = best.get((f"{issuer}.TOTAL_ASSETS", d))
            if got and v.get("FD_TOTAL_ASSETS"):
                ratio = got[1] / v["FD_TOTAL_ASSETS"]
                if not 0.01 < ratio < 100:
                    raise validation.StructuralChangeError(
                        f"kase_ifrs: {issuer} {d} total assets {got[1]:.1f} bn from {Path(got[0]).name} vs {v['FD_TOTAL_ASSETS']:.1f} "
                        f"bn in the KASE key indicators (unit misread?)")
                if abs(ratio - 1) > 0.01:
                    warnings.append(f"{issuer} {d}: total assets {got[1]:.1f} bn vs KASE key indicators {v['FD_TOTAL_ASSETS']:.1f} bn")
            if issuer in ds.get("fin_data", []):
                for line, x in v.items():
                    best[(f"{issuer}.{line}", d)] = (FIN_DATA_URL.format(code=issuer), x)
    raw_store.write_download_manifest("kase", ds["id"], today, {
        "downloaded_at": datetime.now().isoformat(), "source_url": DOCS_URL.format(code="<CODE>"),
        "issuers": list(ds["issuers"]), "fin_data": ds.get("fin_data", []), "files": files})
    records = [{"date": d, "region": dims.NATIONAL, "item_code": code,
                "item_name": f"{ds['issuers'][code.split('.')[0]]}: {names[code.split('.', 1)[1]]}, млрд тенге",
                "value": round(v, 6)} for (code, d), (_, v) in sorted(best.items())]
    got = {r["item_code"].split(".")[0] for r in records if not r["item_code"].split(".", 1)[1].startswith("FD_")}
    if got != set(ds["issuers"]):
        raise validation.StructuralChangeError(f"kase_ifrs: no balance sheet read for {sorted(set(ds['issuers']) - got)}")
    return records, {"frequency": "quarterly", "source_url": DOCS_URL.format(code="<CODE>"),
                     "dataset_id": "KASE issuer financial statements", "note": ds.get("note", ""), "warnings": warnings}


def _order(link: str) -> tuple[int, int]:
    """(year, quarter) of a statement from its file name: …fm1_2024… -> (2024, 1); annual or fm4 -> (year, 4)."""
    m = re.search(r"f(?:m(\d))?_(\d{4})", Path(link).name)
    return (int(m.group(2)), int(m.group(1) or 4)) if m else (0, 0)
