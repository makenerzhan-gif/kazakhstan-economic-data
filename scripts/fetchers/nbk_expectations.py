"""NBK expectation surveys as item-level datasets (added 2026-09-27).

1. INFLATION_EXPECTATIONS_SURVEY -- «Результаты опроса по инфляционным ожиданиям» (xlsx on
   https://nationalbank.kz/ru/page/inflyacionnye-ozhidaniya; the /file/download/<id> changes
   with every monthly edition, so the id is read from the page by the file's title).
   Sheet «Данные»: ~30 questions, each a block «Вопрос №N» / question text / base / a
   «Варианты ответов» row of monthly dates (2016-01 onward) / one row per answer, in % of
   respondents. Item = Q<N>.<answer number in the block> (a question printed twice -- the
   2023/2025 answer-bucket changes of Q6 -- gets .v2 for the second block: Q6.v2.3).
   Sheet «Медианные оценки»: the interval medians the NBK publishes -- MEDIAN.PERCEIVED_12M
   (Q3+Q4), MEDIAN.EXPECTED_12M (Q5+Q6; the scalar INFLATION_EXPECTATIONS series), and
   MEDIAN.EXPECTED_5Y (Q27, from 2025).
   Breaks, from the file's «Комментарии» sheet: wording of Q1-Q6 changed in January 2020
   («продукты питания, непродовольственные товары и услуги» → «товары и услуги»); January
   2022 was not surveyed; from January 2026 Q2, Q16-Q19 were dropped and Q7, Q9, Q10, Q12-Q14,
   Q23-Q25a moved to one month per quarter. A question absent from a month is simply absent.

2. PROFESSIONAL_FORECASTS_SURVEY -- «ХРОНОЛОГИЯ РЕЗУЛЬТАТОВ ОПРОСА» of the NBK
   macroeconomic survey of analysts (https://nationalbank.kz/ru/news/macrosurvey), medians
   per round from August 2022. Date = survey round; item = VARIABLE.<target year> (e.g.
   CPI_DEC.2027 = median forecast of December-on-December inflation in 2027) or VARIABLE for
   the long-run questions (NEUTRAL_RATE, LR_GDP, CPI_5Y).
"""
from __future__ import annotations

import io
import re
import sys
from datetime import date, datetime
from pathlib import Path

import requests

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from fetchers import imf_dims  # noqa: E402
from lib import dims, raw_store, validation  # noqa: E402

BASE = "https://nationalbank.kz"
HEADERS = {"User-Agent": "Mozilla/5.0 (compatible; KZEconDataPipeline/1.0; +https://github.com/)"}
MEDIANS = {"воспринимаемая": ("MEDIAN.PERCEIVED_12M", "Медиана воспринимаемой инфляции за прошедшие 12 месяцев, %"),
           "ожидаемая (в следующие": ("MEDIAN.EXPECTED_12M", "Медиана ожидаемой инфляции в следующие 12 месяцев, %"),
           "ожидаемая (через 5": ("MEDIAN.EXPECTED_5Y", "Медиана ожидаемой инфляции через 5 лет, %")}
FORECAST_VARIABLES = [  # (label prefix, code, has target years)
    ("цена на нефть", "BRENT", True), ("рост ввп", "GDP", True), ("инфляция %, декабрь", "CPI_DEC", True),
    ("базовая ставка", "BASE_RATE", True), ("экспорт", "EXPORTS", True), ("импорт", "IMPORTS", True),
    ("курс usd/kzt", "USDKZT", True), ("нейтральная базовая ставка", "NEUTRAL_RATE", False),
    ("долгосрочный рост ввп", "LR_GDP", False), ("инфляция через 5 лет", "CPI_5Y", False)]


def _clean(text) -> str:
    return " ".join(str(text or "").split())


def _date(cell) -> str | None:
    if isinstance(cell, datetime):
        return cell.date().replace(day=1).isoformat()
    return None


def file_link(page_url: str, title: str) -> str:
    html = requests.get(page_url, headers=HEADERS, timeout=60)
    html.raise_for_status()
    for m in re.finditer(r'href="(/file/download/\d+)"\s*download>\s*([^<]+)<', html.text):
        if _clean(m.group(2)).lower().startswith(title.lower()):
            return BASE + m.group(1)
    raise validation.StructuralChangeError(
        f"STRUCTURAL CHANGE DETECTED in nbk expectations\nWHAT CHANGED: no file titled «{title}» on {page_url}")


def household_records(content: bytes) -> list[dict]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    out: list[dict] = []
    seen: dict[str, int] = {}
    question = text = None
    dates: list[str | None] = []
    answer = 0
    for row in wb["Данные"].iter_rows(values_only=True):
        head = _clean(row[0]) if row else ""
        m = re.match(r"Вопрос\s*№\s*(\d+\s*[а-яa-z]?)$", head)
        if m:
            num = m.group(1).replace(" ", "")
            seen[num] = seen.get(num, 0) + 1
            question = f"Q{num}" + (f".v{seen[num]}" if seen[num] > 1 else "")
            text, dates, answer = _clean(row[1]), [], 0
            continue
        if question is None or len(row) < 3:
            continue
        if _clean(row[1]) == "Варианты ответов":
            dates = [_date(c) for c in row[2:]]
            continue
        if not dates or not _clean(row[1]):
            continue
        answer += 1
        for d, v in zip(dates, row[2:]):
            if d and isinstance(v, (int, float)):
                out.append({"date": d, "region": dims.NATIONAL, "item_code": f"{question}.{answer}",
                            "item_name": f"{question.split('.')[0].replace('Q', 'Вопрос №')}: {text} | {_clean(row[1])}",
                            "value": float(v)})
    header = None
    for row in wb["Медианные оценки"].iter_rows(values_only=True):
        if header is None and any("Воспринимаемая" in _clean(c) for c in row):
            header = [next((v for k, v in MEDIANS.items() if _clean(c).lower().startswith(k)), None) for c in row]
            continue
        d = _date(row[0]) if row else None
        if header and d:
            for spec, v in zip(header, row):
                if spec and isinstance(v, (int, float)):
                    out.append({"date": d, "region": dims.NATIONAL, "item_code": spec[0], "item_name": spec[1], "value": float(v)})
    if header is None or sum(1 for h in header if h) < 2:
        raise validation.StructuralChangeError("nbk expectations: «Медианные оценки» header (Воспринимаемая / Ожидаемая) not found")
    return sorted(out, key=lambda r: (r["item_code"], r["date"]))


def forecaster_records(content: bytes) -> list[dict]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    dates = next(([_date(c) for c in r] for r in rows if sum(1 for c in r if _date(c)) >= 3), None)
    if not dates:
        raise validation.StructuralChangeError("nbk forecasters: no row of survey dates in the chronology")
    first = next(i for i, d in enumerate(dates) if d)
    out, var, unknown = [], None, []
    for r in rows:
        label = _clean(r[first - 2]).lower() if first >= 2 else ""
        if label.startswith("*"):
            continue
        if label:
            hit = next(((code, yearly) for prefix, code, yearly in FORECAST_VARIABLES if label.startswith(prefix)), None)
            if hit is None:
                if any(isinstance(v, (int, float)) for v in r[first:]):
                    unknown.append(label)
                var = None
                continue
            var, name = hit, _clean(r[first - 2])
        if var is None:
            continue
        code, yearly = var
        year = r[first - 1]
        if yearly and not isinstance(year, int):
            continue
        item = f"{code}.{year}" if yearly else code
        for d, v in zip(dates[first:], r[first:]):
            if d and isinstance(v, (int, float)):
                out.append({"date": d, "region": dims.NATIONAL, "item_code": item,
                            "item_name": f"{name}" + (f", прогноз на {year} год" if yearly else "") + " (медиана)",
                            "value": float(v)})
    if unknown:
        raise validation.StructuralChangeError(f"nbk forecasters: unknown variables {unknown} (add to FORECAST_VARIABLES)")
    return sorted(out, key=lambda r: (r["item_code"], r["date"]))


PARSERS = {"households": household_records, "forecasters": forecaster_records}


def fetch(ds: dict) -> tuple[list[dict], dict]:
    fresh = imf_dims.stored_if_fresh(ds, "nbk", ds["id"])
    if fresh:
        return fresh
    url = file_link(ds["page_url"], ds["file_title"])
    resp = requests.get(url, headers=HEADERS, timeout=120)
    resp.raise_for_status()
    today = date.today()
    path = raw_store.save_raw_bytes("nbk", ds["id"], today, "xlsx", resp.content)
    raw_store.write_download_manifest("nbk", ds["id"], today, {
        "downloaded_at": datetime.now().isoformat(), "raw_file": path.name, "source_url": url, "page_url": ds["page_url"]})
    records = PARSERS[ds["survey"]](resp.content)
    if len({r["item_code"] for r in records}) < ds.get("min_items", 5):
        raise validation.StructuralChangeError(f"nbk/{ds['id']}: only {len({r['item_code'] for r in records})} items")
    return records, {"frequency": ds["frequency"], "source_url": url, "dataset_id": ds["file_title"], "note": ds.get("note", "")}
