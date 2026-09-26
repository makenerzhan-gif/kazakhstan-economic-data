"""Income distribution by decile (household budget survey), for CGE/CAEM calibration and
distributional analysis. Added 2026-09-25.

Two BNS sources, reconciled here:

- «Основные показатели дифференциации доходов населения» (stat.gov.kz, living standards,
  listing name=18651): one workbook per edition, quarterly editions from III quarter 2022
  and annual ones (2022-2025) on 2026-09-25. The period is read from the cover («2025 год»,
  «за 2022 год», «I квартал 2026 года») -- the listing's own title/date is not trusted (the
  2022 annual edition is listed as released 24.04.2022, its cover says 24.04.2023). Tables
  are found by title, not by sheet number: «Распределение доходов по 10-процентным группам
  населения» (per decile: income interval, share of income %, mean monthly income per
  capita) and «Основные показатели бедности» (national row: Gini over deciles and over
  quintiles, ratio of the top to the bottom decile).
- Taldau (the same survey): income shares by decile, index 704515, annual 2011-2024 and
  quarterly from 2011; Gini over deciles 704502, annual from 2001 and quarterly from 2011;
  the top/bottom decile ratio 704504, annual 2011-2022; money expenditure of the bottom and
  top decile by item 704518, annual 2001-2024 (the other deciles only 2024).

How they are combined: the editions win wherever they cover a period. Taldau's ANNUAL
series equal the editions exactly (checked 2022 and 2024: all ten shares to the hundredth).
Taldau's QUARTERLY series do not: two quarters are missing (2025-Q4, 2026-Q1), a point is
keyed 2026-06 that no edition confirms, and values differ from the editions in the second
decimal (III quarter 2025, top decile: 24.73 against 24.77). So quarterly Taldau values are
used only BEFORE the first quarterly edition (2011-Q1 to 2022-Q2), and the overlap with the
editions is reported in the note of every dataset.
"""
from __future__ import annotations

import calendar
import csv
import html
import io
import re
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl
import requests

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from lib import dims, raw_store, validation  # noqa: E402

HEADERS = {"User-Agent": "Mozilla/5.0 (compatible; KZEconDataPipeline/1.0; +https://github.com/)"}
LISTING_URL = "https://stat.gov.kz/ru/industries/labor-and-income/stat-life/spreadsheets/?name=18651&year={year}"
FILE_URL = "https://stat.gov.kz/api/iblock/element/{eid}/file/ru/"
TALDAU_URL = "https://taldau.stat.gov.kz/ru/NewIndex/GetIndexTreeData"
FIRST_LISTING_YEAR = 2022
DECILE_TITLE = "Распределение доходов по 10-процентным"
POVERTY_TITLE = "Основные показатели бедности"
QUARTER_RE = re.compile(r"\b(IV|III|II|I)\s*квартал\w*\s+(\d{4})", re.I)
YEAR_RE = re.compile(r"^(?:за\s+)?(\d{4})\s*год", re.I)
QUARTER_NUM = {"i": 1, "ii": 2, "iii": 3, "iv": 4}
DECILE_TERMS = {f"7426{57 + i}": f"D{i + 1:02d}" for i in range(10)}        # 742657 … 742666
DECILE_ROOT = "742656"
_CACHE: dict = {}


# ---------------------------------------------------------------- the BNS editions
def parse_listing(page: str) -> list[str]:
    return [chunk.split('"', 1)[0] for chunk in re.split(r'<div class="divTableRow" id="bx_\d+_', page)[1:]]


def edition_period(cover_rows: list[list]) -> tuple[str, int, int] | None:
    """('annual', year, 12) or ('quarterly', year, quarter) from the cover sheet."""
    for row in cover_rows:
        for cell in row:
            if not isinstance(cell, str):
                continue
            text = " ".join(cell.split())
            q = QUARTER_RE.search(text)
            if q:
                return "quarterly", int(q.group(2)), QUARTER_NUM[q.group(1).lower()]
            y = YEAR_RE.match(text)
            if y:
                return "annual", int(y.group(1)), 12
    return None


def _number(cell) -> float | None:
    if isinstance(cell, (int, float)):
        return float(cell)
    if isinstance(cell, str):
        try:
            return float(cell.replace(" ", "").replace("\xa0", "").replace(",", "."))
        except ValueError:
            return None
    return None


def _find_table(wb, title: str) -> list[list] | None:
    for name in wb.sheetnames:
        rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
        head = " ".join(" ".join(str(c).split()) for r in rows[:4] for c in r if isinstance(c, str))
        if title in head:
            return rows
    return None


def parse_decile_table(rows: list[list]) -> dict[str, dict[str, float]]:
    """{D01: {share, mean, lower, upper}} from «Распределение доходов по 10-процентным группам»."""
    hi = next((i for i, r in enumerate(rows) if any(isinstance(c, str) and "Границы" in c for c in r)), None)
    if hi is None:
        return {}
    head = rows[hi]
    col = {}
    for j, c in enumerate(head):
        t = " ".join(str(c).split()) if isinstance(c, str) else ""
        if t.startswith("Границы"):
            col["bounds"] = j
        elif t.startswith("Доля"):
            col["share"] = j
        elif t.startswith("Среднемесячный"):
            col["mean"] = j
    if set(col) != {"bounds", "share", "mean"}:
        return {}
    out = {}
    for r in rows[hi + 1:]:
        label = _number(r[0]) if r else None
        if label is None or not 1 <= label <= 10 or label != int(label):
            continue
        bounds = str(r[col["bounds"]] or "").replace(" ", "").replace("\xa0", "")
        m = re.fullmatch(r"(\d+)-(\d+)", bounds)
        share, mean = _number(r[col["share"]]), _number(r[col["mean"]])
        if m is None or share is None or mean is None:
            continue
        out[f"D{int(label):02d}"] = {"share": share, "mean": mean, "lower": float(m.group(1)), "upper": float(m.group(2))}
    return out


def parse_poverty_table(rows: list[list]) -> dict[str, float]:
    """National row of «Основные показатели бедности»: depth, severity, gini10, gini20, ratio."""
    hi = next((i for i, r in enumerate(rows) if any(isinstance(c, str) and "Джини" in c for c in r)), None)
    if hi is None:
        return {}
    col = {}
    for j, c in enumerate(rows[hi]):
        t = " ".join(str(c).split()) if isinstance(c, str) else ""
        if t.startswith("Глубина"):
            col["depth"] = j
        elif t.startswith("Острота"):
            col["severity"] = j
        elif "Джини" in t and "10%" in t:
            col["gini10"] = j
        elif "Джини" in t and "20%" in t:
            col["gini20"] = j
        elif t.startswith("Соотношение") or (t == "" and j == 5 and "gini20" in col):
            col["ratio"] = j                  # the 2022 annual edition leaves this header cell blank
    for r in rows[hi + 1:]:
        label = " ".join(str(r[0]).split()) if r and isinstance(r[0], str) else ""
        if label.startswith(("Қазақстан Республикасы", "Республика Казахстан")):
            return {k: v for k, j in col.items() if (v := _number(r[j])) is not None}
    return {}


def editions() -> list[dict]:
    """Every edition on the listing: {eid, kind, year, sub, deciles, poverty}."""
    if "editions" in _CACHE:
        return _CACHE["editions"]
    eids: list[str] = []
    for year in range(FIRST_LISTING_YEAR, date.today().year + 1):
        page = requests.get(LISTING_URL.format(year=year), headers=HEADERS, timeout=60).text
        eids += [e for e in parse_listing(page) if e not in eids]
    out, today = [], date.today()
    for eid in eids:
        content = requests.get(FILE_URL.format(eid=eid), headers=HEADERS, timeout=120).content
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception:  # noqa: BLE001 -- a non-workbook entry must not sink the others
            continue
        cover = next((s for s in wb.sheetnames if s.strip().startswith("Обложка")), wb.sheetnames[0])
        period = edition_period([list(r) for r in wb[cover].iter_rows(max_row=30, values_only=True)])
        dec = _find_table(wb, DECILE_TITLE)
        pov = _find_table(wb, POVERTY_TITLE)
        if period is None or dec is None:
            continue
        rec = {"eid": eid, "kind": period[0], "year": period[1], "sub": period[2],
               "deciles": parse_decile_table(dec), "poverty": parse_poverty_table(pov) if pov else {}}
        if len(rec["deciles"]) != 10:
            raise validation.StructuralChangeError(
                f"bns/18651 edition {eid}: the decile table parsed to {len(rec['deciles'])} rows, expected 10")
        path = raw_store.save_raw_bytes("bns", f"INCOME_DIFFERENTIATION_{eid}", today, "xlsx", content)
        raw_store.write_download_manifest("bns", f"INCOME_DIFFERENTIATION_{eid}", today, {
            "downloaded_at": datetime.now().isoformat(), "raw_file": path.name, "source_url": FILE_URL.format(eid=eid),
            "period": f"{period[1]}{'Q' + str(period[2]) if period[0] == 'quarterly' else ''}"})
        out.append(rec)
    keys = [(e["kind"], e["year"], e["sub"]) for e in out]
    dupes = {k for k in keys if keys.count(k) > 1}
    if dupes:
        raise validation.StructuralChangeError(f"bns/18651: several editions for the same period: {sorted(dupes)}")
    _CACHE["editions"] = out
    return out


def period_date(kind: str, year: int, sub: int) -> str:
    return f"{year}-12-31" if kind == "annual" else f"{year}-{3 * sub - 2:02d}-01"


# ---------------------------------------------------------------- Taldau history
def _taldau(index: str, period: str, measure: str, terms: str, dics: str, parent: str = "", idx: str = "0",
            term: str = "741880") -> list[dict]:
    body = {"p_parent_id": parent, "p_index_id": index, "p_keyword": "", "p_period_id": period, "p_measure_id": measure,
            "p_term_id": term, "p_terms": terms, "p_dicIds": dics, "idx": idx,
            "filter": '[{"property":null,"value":null}]', "id": ""}
    resp = requests.post(TALDAU_URL, data=body, headers={**HEADERS, "X-Requested-With": "XMLHttpRequest"}, timeout=60)
    resp.raise_for_status()
    today = date.today()
    name = f"TALDAU_{index}_{period}_{(parent or 'root')}_{terms.replace(',', '-')}"[:120]
    raw_store.save_raw_bytes("bns", name, today, "json", resp.content)
    return resp.json()


def taldau_values(node: dict, kind: str) -> dict[str, float]:
    out = {}
    for k, v in node.items():
        if not (isinstance(k, str) and len(k) == 7 and k[0] == "y" and k[1:].isdigit()) or v in (None, ""):
            continue
        month, year = int(k[1:3]), int(k[3:])
        value = _number(v)
        if value is None:
            continue
        out[period_date(kind, year, 12 if kind == "annual" else (month + 2) // 3)] = value
    return out


def taldau_decile_shares(kind: str) -> dict[str, dict[str, float]]:
    """{D01: {date: share}} from Taldau 704515."""
    nodes = _taldau("704515", "7" if kind == "annual" else "5", "7", f"741880,{DECILE_ROOT}", "67,838",
                    parent=DECILE_ROOT, idx="1", term=DECILE_ROOT)
    return {DECILE_TERMS[str(n["id"])]: taldau_values(n, kind) for n in nodes if str(n.get("id")) in DECILE_TERMS}


def taldau_scalar(index: str, kind: str, measure: str = "758") -> dict[str, float]:
    nodes = _taldau(index, "7" if kind == "annual" else "5", measure, "741880", "67")
    return taldau_values(nodes[0], kind) if nodes else {}


# ---------------------------------------------------------------- datasets
def _structural(what: str) -> None:
    raise validation.StructuralChangeError(f"STRUCTURAL CHANGE DETECTED in bns income distribution: {what}")


def _decile_records(values: dict[str, dict[str, float]], label: str) -> list[dict]:
    return [{"date": d, "region": dims.NATIONAL, "item_code": code, "item_name": f"{int(code[1:])}-я децильная группа ({label})",
             "value": v} for code, by_date in sorted(values.items()) for d, v in sorted(by_date.items())]


def fetch_deciles(ds: dict) -> tuple[list[dict], dict]:
    kind, measure = ds["kind"], ds["measure"]           # kind: annual | quarterly; measure: share | mean | upper
    eds = [e for e in editions() if e["kind"] == kind]
    if not eds:
        _structural(f"no {kind} edition of «{DECILE_TITLE}» on {LISTING_URL}")
    values: dict[str, dict[str, float]] = {}
    for e in eds:
        for code, row in e["deciles"].items():
            values.setdefault(code, {})[period_date(kind, e["year"], e["sub"])] = row[measure]
    note = ds.get("note", "")
    if measure == "share":
        hist = taldau_decile_shares(kind)
        first_edition = min(period_date(kind, e["year"], e["sub"]) for e in eds)
        diffs, overlap = [], 0
        for code, by_date in hist.items():
            for d, v in by_date.items():
                if d in values.get(code, {}):
                    overlap += 1
                    if abs(values[code][d] - v) > 0.005:
                        diffs.append((d, code, v, values[code][d]))
                elif kind == "annual" or d < first_edition:
                    values.setdefault(code, {})[d] = v
        if kind == "annual" and diffs:
            _structural(f"Taldau 704515 annual differs from the BNS editions: {diffs[:5]}")
        note += (f" Taldau 704515 fills the periods no edition covers{' (before ' + first_edition + ')' if kind == 'quarterly' else ''};"
                 f" on the {overlap} overlapping points it differs from the editions in {len(diffs)}"
                 + (f", e.g. {diffs[0][0]} {diffs[0][1]}: {diffs[0][2]} in Taldau against {diffs[0][3]}" if diffs else "")
                 + " -- the editions are kept.")
    for d in {d for by in values.values() for d in by}:
        if measure == "share":
            total = sum(by.get(d, 0.0) for by in values.values())
            if len([1 for by in values.values() if d in by]) == 10 and abs(total - 100) > 0.1:
                _structural(f"decile shares for {d} sum to {total:.2f}")
    label = {"share": "доля доходов, %", "mean": "среднемесячный доход на душу, тенге",
             "upper": "верхняя граница интервала, тенге в месяц"}[measure]
    return _decile_records(values, label), {
        "frequency": kind, "source_url": LISTING_URL.format(year="YYYY"), "dataset_id": "bns/18651" + (
            "+taldau/704515" if measure == "share" else ""),
        "note": note + f" Editions read: {', '.join(e['eid'] + ' ' + str(e['year']) + ('Q' + str(e['sub']) if kind == 'quarterly' else '') for e in sorted(eds, key=lambda e: (e['year'], e['sub'])))}."}


EXPENDITURE_ITEMS = {"545138": "MONEY_EXPENDITURE", "545139": "CONSUMER_EXPENDITURE", "545140": "FOOD",
                     "545142": "EATING_OUT", "545144": "NONFOOD", "545145": "SERVICES",
                     "545199": "TRANSFERS_TO_HOUSEHOLDS", "545205": "TAXES_AND_PAYMENTS", "545210": "DEBT_REPAYMENT"}


def fetch_expenditure_tails(ds: dict) -> tuple[list[dict], dict]:
    """Money expenditure of the bottom (D01) and top (D10) decile by item, Taldau 704518."""
    records = []
    for item, name in EXPENDITURE_ITEMS.items():
        nodes = _taldau("704518", "7", "1", f"741880,741917,808297,{item},{DECILE_ROOT}", "67,749,270,735,838",
                        parent=DECILE_ROOT, idx="4", term=DECILE_ROOT)
        for n in nodes:
            code = DECILE_TERMS.get(str(n.get("id")))
            if code not in ("D01", "D10"):
                continue                                   # the middle deciles exist for 2024 only
            for d, v in taldau_values(n, "annual").items():
                records.append({"date": d, "region": dims.NATIONAL, "item_code": f"{code}_{name}",
                                "item_name": f"{code}: {name.replace('_', ' ').lower()}", "value": v})
    if len({r["item_code"] for r in records}) < 12:
        _structural(f"Taldau 704518 gave only {len({r['item_code'] for r in records})} decile x item series")
    return records, {"frequency": "annual", "source_url": TALDAU_URL, "dataset_id": "taldau/704518",
                     "note": ds.get("note", "")}


def fetch(ds: dict) -> tuple[list[dict], dict]:
    return {"deciles": fetch_deciles, "expenditure_tails": fetch_expenditure_tails}[ds["table"]](ds)


# ---------------------------------------------------------------- scalar series (update_bns)
def _load_processed(indicator_id: str) -> dict[str, float]:
    path = Path(__file__).resolve().parents[2] / "data" / "processed" / "bns" / f"{indicator_id.lower()}.csv"
    if not path.exists():
        return {}
    with path.open(encoding="utf-8") as f:
        return {r["date"]: float(r["value"]) for r in csv.DictReader(f) if r.get("value") not in (None, "")}


def poverty_series(indicator_id: str, key: str, kind: str, taldau_index: str | None, note: str) -> tuple[list[dict], dict]:
    """A national row of «Основные показатели бедности» from every edition, earlier periods from
    Taldau (annual: wherever no edition; quarterly: before the first quarterly edition only)."""
    eds = [e for e in editions() if e["kind"] == kind and key in e["poverty"]]
    values = {period_date(kind, e["year"], e["sub"]): e["poverty"][key] for e in eds}
    first = min(values) if values else "9999"
    diffs = []
    if taldau_index:
        for d, v in taldau_scalar(taldau_index, kind).items():
            if d in values:
                if abs(values[d] - v) > 0.0015:
                    diffs.append((d, v, values[d]))
            elif kind == "annual" or d < first:
                values[d] = v
    for d, v in _load_processed(indicator_id).items():        # history no source still serves
        values.setdefault(d, v)
    if not values:
        _structural(f"{indicator_id}: no edition and no Taldau history")
    return [{"date": d, "value": v} for d, v in sorted(values.items())], {
        "frequency": kind, "source_url": LISTING_URL.format(year="YYYY"),
        "dataset_id": "bns/18651 «Основные показатели бедности»" + (f" + taldau/{taldau_index}" if taldau_index else ""),
        "note": note + f" {len(eds)} editions" + (f"; Taldau {taldau_index} for the earlier periods, differing from the "
                                                   f"editions on {len(diffs)} overlapping points" if taldau_index else "") + "."}
