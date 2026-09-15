"""BNS external trade by commodity group — the dimensional view of the export workbook
that already feeds EXPORTS and OIL_EXPORTS_* (element 446905, see the module comment
above `_parse_trade_workbook_by_hs` in fetchers/bns.py for the file's structure).

Groups are HS (ТН ВЭД ЕАЭС) prefixes listed in dictionaries/hs_export_groups.csv — a
4-digit heading (1001 wheat), a chapter (15 fats and oils), a 6-digit line (271121
natural gas) or several of them (7208–7212 flat-rolled). The workbook is a flat
6-digit breakdown inside regional blocks with no national product rows, so a group is
the sum of its 6-digit lines over the regions. That sum is legitimate only while the
regional product rows reproduce the published national total, which is re-checked for
every month here exactly as fetchers/bns.py does: a month that fails the check is
skipped and reported, a broad failure stops the dataset. TOTAL is the published
national row itself, never a sum.

Two datasets read the same file: value (thousand USD) and volume (tonnes), parsed once
per process. The scalar EXPORTS, OIL_EXPORTS_VALUE and OIL_EXPORTS_VOLUME series are
derived from these datasets (config/indicators.yaml `derived_from`, scripts/
update_derived.py) — the workbook is downloaded once per run and archived once per day
(raw id EXPORTS, the name the file has always had in data/raw/bns).
"""
from __future__ import annotations

import io
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from lib import dims, raw_store, validation  # noqa: E402
from fetchers import bns  # noqa: E402

ELEMENT_ID = 446905
URL = f"https://stat.gov.kz/api/iblock/element/{ELEMENT_ID}/file/ru/"
RAW_ID = "EXPORTS"
UNIT_OFFSET = {"tonnes": 0, "usd": 2}      # each month spans [tonnes, additional unit, thousand USD]
_PARSED: dict[tuple[str, str], dict] = {}


def _structural(ds: dict, what: str, expected, actual) -> None:
    raise validation.StructuralChangeError("\n".join([
        f"STRUCTURAL CHANGE DETECTED in bns/{ds['id']} (element {ELEMENT_ID})",
        f"WHAT CHANGED: {what}", f"EXPECTED: {expected}", f"ACTUAL: {actual}",
        "ACTION REQUIRED: open the export workbook, compare with the layout described in "
        "scripts/fetchers/bns_trade.py and fetchers/bns.py, update the parser or "
        "dictionaries/hs_export_groups.csv, then re-run.",
    ]))


def load_groups(name: str) -> list[dict]:
    """Dictionary entries with their HS prefix lists ('7208|7209' -> ['7208', '7209'])."""
    out = []
    for e in dims.load_dictionary(name):
        out.append({**e, "prefixes": [p.strip() for p in (e.get("hs") or "").split("|") if p.strip()]})
    return out


def _content(today: date) -> tuple[bytes, str]:
    """(workbook bytes, the raw file holding them). bns._download caches per process,
    raw_store stores identical bytes once per day."""
    content = bns._download(URL)
    path = raw_store.save_raw_bytes("bns", RAW_ID, today, "xlsx", content)
    raw_store.write_download_manifest("bns", RAW_ID, today, {
        "downloaded_at": datetime.now().isoformat(), "source_url": URL, "element_id": ELEMENT_ID, "raw_file": path.name,
        "note": "read by EXPORTS_VALUE_BY_COMMODITY_GROUP and EXPORTS_VOLUME_BY_COMMODITY_GROUP; "
                "EXPORTS, OIL_EXPORTS_VALUE and OIL_EXPORTS_VOLUME are derived from them"})
    return content, path.name


def parse_groups(content: bytes, groups: list[dict], ds: dict) -> dict[str, dict[str, dict[str, float]]]:
    """{unit: {iso month: {group code: value}}} for 'tonnes' and 'usd', every sheet of the
    workbook; months whose regional rows do not reproduce the national total are skipped."""
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    out: dict[str, dict[str, dict[str, float]]] = {"tonnes": {}, "usd": {}}
    skipped: list[tuple] = []
    checked = 0
    by_prefix = [(p, g["code"]) for g in groups for p in g["prefixes"]]
    for sheet_name in wb.sheetnames:
        if sheet_name in ("Метаданные", "Показатель"):
            continue
        ws = wb[sheet_name]
        header_row = national_row = None
        sums: dict[str, dict[int, float]] = defaultdict(lambda: defaultdict(float))
        all_sums: dict[int, float] = defaultdict(float)
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 1:
                header_row = row
                continue
            if i == 3:
                national_row = row
                label = (row[0] or "").strip() if row and row[0] else ""
                if label != bns.TRADE_TOTAL_ROW_LABEL:
                    _structural(ds, f"row 4 of sheet {sheet_name!r} is not the national total", bns.TRADE_TOTAL_ROW_LABEL, label)
                continue
            if i < 4 or not row or row[0] is None:
                continue
            code = str(row[0]).strip()
            if not code.isdigit():
                continue                                  # a region heading
            if len(code) != 6:
                _structural(ds, f"sheet {sheet_name!r} row {i + 1} has a {len(code)}-digit code {code!r}",
                            "a flat 6-digit breakdown (mixed lengths would double-count)", code)
            targets = [g for p, g in by_prefix if code.startswith(p)]
            for col, value in enumerate(row):
                if not isinstance(value, (int, float)) or isinstance(value, bool):
                    continue
                all_sums[col] += value
                for g in targets:
                    sums[g][col] += value
        if header_row is None or national_row is None:
            _structural(ds, f"sheet {sheet_name!r} has no header/national rows", "header at row 2, national total at row 4", sheet_name)
        for col_idx, cell in enumerate(header_row):
            m = bns.MONTH_HEADER_RE.match(str(cell).strip()) if cell else None
            month = bns.RU_MONTHS.get(m.group(1).lower()) if m else None
            if not month:
                continue
            iso = f"{int(m.group(2)):04d}-{month:02d}-01"
            for unit, offset in UNIT_OFFSET.items():
                col = col_idx + offset
                national = national_row[col] if col < len(national_row) else None
                if not isinstance(national, (int, float)) or not national:
                    continue
                checked += 1
                if abs(all_sums.get(col, 0.0) - national) > abs(national) * 1e-6:
                    # The group sums are unverifiable for this month and are withheld; the
                    # published national row is not a sum and stays (EXPORTS is derived from it
                    # and must be gap-free).
                    skipped.append((sheet_name, m.group(0), unit, round((all_sums.get(col, 0.0) / national - 1) * 100, 4)))
                    out[unit][iso] = {"TOTAL": float(national)}
                    continue
                month_vals = {g["code"]: sums[g["code"]].get(col, 0.0) for g in groups if g["prefixes"]}
                month_vals["TOTAL"] = float(national)
                out[unit][iso] = month_vals
    if checked and len(skipped) > checked * 0.1:
        _structural(ds, f"regional product rows failed to sum to the national total in {len(skipped)} of {checked} month/unit checks",
                    "at most a couple of isolated months (2 of 180 as of 2026-09-01)", skipped[:8])
    if skipped:
        print(f"bns/{ds['id']}: skipped {len(skipped)} unverifiable month(s): {skipped}", file=sys.stderr)
    if not out["usd"]:
        _structural(ds, "no month passed the partition check", "monthly national totals matched by the regional rows", sheet_name)
    return out


def fetch(ds: dict) -> tuple[list[dict], dict]:
    """ds['measure'] is 'usd' (thousand USD) or 'tonnes'."""
    today = date.today()
    content, origin = _content(today)
    groups = load_groups(ds["dictionary"])
    key = (URL, ds["dictionary"])
    if key not in _PARSED:
        _PARSED[key] = parse_groups(content, groups, ds)
    by_month = _PARSED[key][ds["measure"]]
    names = {g["code"]: g["name_ru"] for g in groups}
    records = [{"date": iso, "region": dims.NATIONAL, "item_code": code, "item_name": names[code], "value": v}
               for iso in sorted(by_month) for code, v in by_month[iso].items()]
    return records, {"frequency": "monthly", "source_url": URL, "dataset_id": f"{ELEMENT_ID},measure={ds['measure']}",
                     "note": ds.get("note", "") + f" Raw file: {origin}.", "warnings": [], "raw_file": origin}
