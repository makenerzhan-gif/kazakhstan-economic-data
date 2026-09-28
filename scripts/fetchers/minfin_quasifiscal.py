"""Quasi-fiscal flows through the republican budget, January-to-month, from Minfin's monthly
Statistical Bulletin (added 2026-09-27): budget loans to, and equity injected into, the national
holdings and development institutions, by recipient.

Sheet «табл 8 (расх)» («табл 8(расходы)» / «табл 8(расх)» in 2013-2014) of every bulletin edition,
read through fetchers/minfin (so the archive mode and the backfill cache of
scripts/backfill_minfin_bulletins.py apply, as for fetchers/minfin_regions.py). It is the republican
budget's spending by functional group / administrator (АБП) / programme, in two sections that
matter here:
  * «БЮДЖЕТНЫЕ КРЕДИТЫ» -> LOAN, budget loans;
  * «ПРИОБРЕТЕНИЕ ФИНАНСОВЫХ АКТИВОВ» -> EQUITY, acquisition of financial assets (mostly
    charter-capital increases).
The layout is the same 7 columns since 2013: A function, B subfunction, C administrator,
D programme, E Kazakh name, F value (million KZT, January to the month), G Russian name. A
programme row has a code in D; the administrator code in C is printed on the administrator's own
row and carried forward.

Recipients are read from the Russian programme label, not the codes: administrator and programme
codes change with every government reorganisation (Baiterek -> DBK: 242/217 -> 249/036 -> 229/036;
Отбасы 242/231 -> 249/231 -> 229/231; КазАгро 212/023 (2013-2016) -> ACC 212/262). The first entity
a label names is the direct counterparty (usually the holding: «Кредитование АО «НУХ «Байтерек» с
последующим кредитованием АО «Банк Развития Казахстана» …»); the last one named before «через» is
the final recipient («… АО «Фонд развития промышленности» через АО «Банк Развития Казахстана»»
makes DBK a conduit). АО «БРК-Лизинг» became АО «Фонд развития промышленности» in 2020 (programmes
249/218 and 249/243 change name between the 2020-09 and 2020-10 editions): both are IDF.

Items (all million KZT, year to date, dated the first day of the period's last month:
2026-07-01 = January-July 2026), for SEC in LOAN, EQUITY:
  SEC                  section total (the section header row);
  SEC.DI               development institutions and national holdings (the sum of the next);
  SEC.<X>              by final recipient, X in BAITEREK (the holding itself), DBK, IDF, DAMU,
                       KAZAKHEXPORT, QIC, KAZAGRO, ACC, KAF, KAZAGROGARANT, FCC (Продкорпорация),
                       OTBASY, KHC, KMC, HCGF, KMGF, SAMRUK, KTZ, FPK (Фонд проблемных кредитов);
  SEC.FAM.<F>          by economic family of the final recipient, F in DEV_FINANCE, SME, AGRI,
                       HOUSING, SAMRUK, FPK -- stable across the 2021 move of KazAgro's companies
                       into Baiterek;
  SEC.VIA.<H>          by direct counterparty, H in BAITEREK, KAZAGRO, SAMRUK (what the holding
                       received, for itself or to pass on);
  SEC.LOCAL_GOV, SEC.GUARANTEE_CALL, SEC.INTL_ORG, SEC.OTHER_QUASI, SEC.OTHER
                       the rest, so that SEC = SEC.DI + these five: loans to local executive
                       bodies, payments under state guarantees, shares of international
                       organisations (МФО, ЕАБР, Тюркский инвестфонд), other state companies'
                       charter capital (АО/НАО/РГП/ТОО), and anything else.
Every edition carries every item (0 when no programme of it is listed: January editions often
have no loan or asset rows at all); items that are zero in every edition are dropped.

Checks (fetch):
  * hard: the programme rows of a section add up to its header (0.5 mln): 161/161 loan and
    160/160 asset sections of the 2013-01 … 2026-07 editions (the 2013-01 edition has no asset
    section);
  * warning: the section headers equal the «Бюджетные кредиты» / «Приобретение финансовых
    активов» lines of табл 7 in the column of the same period (December: the year's report
    column). Known source errors: табл 7 of the 2021-02 and 2021-03 editions shows a stale
    80,337.874 against table 8's 88,237.874 and 96,400.662;
  * warning: LOAN.DI = табл 10 economic specifics 513 + 519 (exact in every edition with both).
    табл 10's 612 is NOT a check for EQUITY: it also codes charter capital financed through
    targeted transfers to regions (12-118 bn more in most years);
  * warning: every item's year-to-date value is non-decreasing within a year;
  * warning: a loan label naming a company that no recipient pattern knows, and a label cut at
    ~240 characters whose final recipient is not printed (KNOWN_TRUNCATED resolves the known ones).

Verified (prototype over all 161 cached editions, reproduced by this module): December YTD, bn KZT,
2013 -> 2025: LOAN 122.1, 118.5, 190.8, 315.0, 282.5, 250.3, 381.6, 338.2, 389.6, 655.6, 661.4,
437.8, 442.1; LOAN.DI 60.0, 58.9, 68.9, 182.5, 194.1, 146.5, 246.5, 245.2, 284.7, 577.3, 410.3,
274.0, 235.5; EQUITY.DI 98.5, 433.1 (of which FPK 250.0), 136.8, 231.3, 133.4, 17.9, 126.0, 177.6,
14.9, 70.0, 204.5, 40.0, 30.0. Jan-Dec 2020: LOAN 338,215.473 of which DI 245,200 (ACC 70,000,
Baiterek->DBK 70,000, ->IDF 23,700, Baiterek 22,500, Отбасы 59,000).

Pitfalls:
  * Edition period = the latest January-to-month header among табл 7, 8 (расх) and 10: табл 8's
    own header is stale in the 2018-10 and 2020-07 editions whose numbers are current; headers
    without «отчет» in 2018-03 and 2019-10 («январь-март 2018 г.»).
  * The January 2014 edition appends a second (Jan-Dec 2013) block: reading stops at the first
    repeated section.
  * Column A is misaligned in the 2018-08 edition (a stray «10» on the asset header row): section
    headers are recognised by label with no administrator/programme code.
  * Values may be strings («  677 000», 2017); a «103 » code prefixes some 2020 labels.
  * «космического ракетного комплекса «Байтерек»» is not the holding (the holding pattern needs
    «холдинг» or «НУХ»).
  * gov.kz titles are unreliable: the Jan-Dec 2025 edition (document 964016) is titled only «(Is
    developed in Kazakh and Russian …)» and no title search finds it (UNTITLED_BULLETINS).
  * Missing editions: 2020-04 and 2026-04 are not on gov.kz; the next month's flow then covers two.
  * Out of scope: quasi-fiscal current transfers booked in the ЗАТРАТЫ section (217/202 «Целевое
    перечисление в АО «Фонд проблемных кредитов»» 2,092.9 bn in 2017; Даму subsidies 243/087),
    National Fund / NBK / UAPF purchases of holding bonds, local-budget lending (in
    REGIONAL_BUDGETS_YTD NET_LENDING).
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from fetchers import minfin  # noqa: E402
from lib import dims, validation  # noqa: E402

DATASET_ID = "RB_QUASIFISCAL_YTD"
TABLE8, TABLE7, TABLE10 = "таб8(расх)", "таб7", "таб10"
TOLERANCE = 0.5          # million KZT
# gov.kz documents that are bulletin editions although their title does not say so (found by
# opening every xlsx of the budget direction's listing, 2026-09-27)
UNTITLED_BULLETINS = {"964016": "January-December 2025"}
# табл 7 errors of the source (the section totals of табл 8 are right: their programme rows add up)
KNOWN_TABLE7_ERRORS = {"2021-02-01", "2021-03-01"}

SECTIONS = {"LOAN": "Бюджетные кредиты", "EQUITY": "Приобретение финансовых активов"}
SECTION_RE = {
    "LOAN": re.compile(r"^бюджетные кредиты$"),
    "EQUITY": re.compile(r"^приобретение финансовых активов$"),
    "EXP": re.compile(r"^(ii\.\s*)?затраты$"),
    "REPAY": re.compile(r"^погашение займов$"),
    "TOTAL": re.compile(r"^расходы$"),
}

# (code, name, regex on the normalised Russian label). The holdings first: the first entity a
# label names is the direct counterparty, the last one before «через» the final recipient.
ENTITIES = [
    ("BAITEREK", "АО «НУХ «Байтерек»", r"холдинг\w*\s*[«\"]?\s*байтерек|нух\s*[«\"]?\s*байтерек|холдинг\s+байтерек"),
    ("KAZAGRO", "АО «НУХ «КазАгро»", r"холдинг\w*\s*[«\"]?\s*казагро\b"),
    ("SAMRUK", "АО «ФНБ «Самрук-Казына»", r"самрук-?\s*қазына|самрук-?\s*казына"),
    ("DBK_LEASING", "АО «БРК-Лизинг»", r"брк-?\s*лизинг"),
    ("DBK", "АО «Банк Развития Казахстана»", r"банк\w*\s*развития\s+казахстана"),
    ("DAMU", "АО «Фонд развития предпринимательства «Даму»", r"фонд\w*\s+развития\s+предпринимательства\s*[«\"]?\s*даму"),
    ("ACC", "АО «Аграрная кредитная корпорация»", r"аграрн\w+\s+кредитн\w+\s+корпораци"),
    ("KAF", "АО «КазАгроФинанс»", r"казагрофинанс"),
    ("OTBASY", "АО «Отбасы банк» (ЖССБК)", r"жилищн\w+\s+строительн\w+\s+сберегательн\w+\s+банк|отбасы\s+банк"),
    ("KHC", "АО «Казахстанская жилищная компания»", r"казахстанск\w+\s+жилищн\w+\s+компани"),
    ("KMC", "АО «Казахстанская ипотечная компания»", r"казахстанск\w+\s+ипотечн\w+\s+компани"),
    ("HCGF", "АО «Казахстанский фонд гарантирования жилищного строительства»", r"фонд\w*\s+гарантирования\s+жилищного\s+строительства"),
    ("KMGF", "АО «Казахстанский фонд гарантирования ипотечных кредитов»", r"фонд\w*\s+гарантирования\s+ипотечных\s+кредитов"),
    ("IDF", "АО «Фонд развития промышленности» (до 2020 «БРК-Лизинг»)", r"фонд\w*\s+развития\s+промышленности"),
    ("KAZAKHEXPORT", "АО «ЭКА «KazakhExport»", r"kazakhexport|экспортн\w+(\s*-\s*|\s+)кредитн\w+|экспортн\w+\s+страхов\w+"),
    ("QIC", "АО «Qazaqstan Investment Corporation» (Казына Капитал Менеджмент)", r"qazaqstan\s+investment\s+corporation|казына\s+капитал\s+менеджмент"),
    ("FPK", "АО «Фонд проблемных кредитов»", r"фонд\w*\s+проблемных\s+кредитов"),
    ("KTZ", "АО «НК «Қазақстан темір жолы»", r"қазақстан\s+темір\s+жолы|казакстан\s+темир\s+жолы"),
    ("FCC", "АО «НК «Продовольственная контрактная корпорация»", r"продовольственн\w+\s+контрактн\w+\s+корпораци"),
    ("KAZAGROGARANT", "АО «Фонд гарантирования исполнения обязательств» (КазАгроГарант)", r"фонд\w*\s+гарантирования\s+исполнения\s+обязательств|казагрогарант"),
]
_ENTITIES = [(c, re.compile(p)) for c, _, p in ENTITIES]
RENAME = {"DBK_LEASING": "IDF"}
RECIPIENTS = [c for c, _, _ in ENTITIES if c not in RENAME]
RECIPIENT_NAMES = {c: n for c, n, _ in ENTITIES}
HOLDINGS = ("BAITEREK", "KAZAGRO", "SAMRUK")
FAMILY = {
    "BAITEREK": "DEV_FINANCE", "DBK": "DEV_FINANCE", "IDF": "DEV_FINANCE", "QIC": "DEV_FINANCE", "KAZAKHEXPORT": "DEV_FINANCE",
    "DAMU": "SME",
    "KAZAGRO": "AGRI", "ACC": "AGRI", "KAF": "AGRI", "KAZAGROGARANT": "AGRI", "FCC": "AGRI",
    "OTBASY": "HOUSING", "KHC": "HOUSING", "KMC": "HOUSING", "HCGF": "HOUSING", "KMGF": "HOUSING",
    "SAMRUK": "SAMRUK", "KTZ": "SAMRUK",
    "FPK": "FPK"}
FAMILY_NAMES = {"DEV_FINANCE": "развитие (Байтерек, БРК, ФРП, KazakhExport, QIC)", "SME": "МСБ (Даму)",
                "AGRI": "АПК (КазАгро, АКК, КазАгроФинанс, Продкорпорация)",
                "HOUSING": "жильё (Отбасы банк, КЖК, КИК, фонды гарантирования)",
                "SAMRUK": "Самрук-Казына и КТЖ", "FPK": "Фонд проблемных кредитов"}

# recipient class of a programme that names no development institution
OTHER_RULES = [
    ("LOCAL_GOV", "местные исполнительные органы", r"областн\w+\s+бюджет|местн\w+\s+исполнительн\w+\s+орган|бюджетов\s+город"),
    ("GUARANTEE_CALL", "выполнение обязательств по государственным гарантиям", r"обязательств\w*\s+по\s+государственн\w+\s+гаранти"),
    ("INTL_ORG", "международные организации", r"международн\w+\s+(финансов\w+\s+)?организаци|евразийск\w+\s+банк|тюркск\w+\s+инвестиционн"
     r"|азиатск\w+\s+банк|исламск\w+\s+банк|международн\w+\s+банк"),
    ("OTHER_QUASI", "прочие государственные компании (АО/НАО/РГП/ТОО)", r"уставн\w+\s+капитал|акционерн\w+\s+обществ|\bао\b|\bнао\b|\bргп\b|\bтоо\b"),
]
_OTHER_RULES = [(c, re.compile(p)) for c, _, p in OTHER_RULES]
REMAINDERS = [c for c, _, _ in OTHER_RULES] + ["OTHER"]
REMAINDER_NAMES = {**{c: n for c, n, _ in OTHER_RULES}, "OTHER": "прочее"}
# loan labels that name a company no ENTITIES pattern knows, classified OTHER_QUASI on purpose
KNOWN_OTHER_LOANS = (r"центр\w*\s+модернизации\s+и\s+развития\s+жилищно-коммунального",)

# labels cut at ~240 characters before the final recipient: (section, administrator, programme,
# year) -> final recipient, from the full label in Minfin's «Отчет об исполнении республиканского
# бюджета». None = the last entity printed is the final recipient (КТЖ passes the loan on to its
# own subsidiary АО «Пассажирские перевозки»).
KNOWN_TRUNCATED = {("EQUITY", 229, 7, 2026): "IDF",        # Байтерек -> БРК -> Фонд развития промышленности
                   ("LOAN", 228, 4, 2025): None, ("LOAN", 228, 4, 2026): None}
TRUNCATED_AT = 230
VIA_RE = re.compile(r"\bчерез\b")
NEXT_RE = re.compile(r"с\s+последующ")


def item_names() -> dict[str, str]:
    """item_code -> Russian item name, in output order."""
    out = {}
    for sec, title in SECTIONS.items():
        out[sec] = f"{title} (республиканский бюджет), всего"
        out[f"{sec}.DI"] = f"{title}: институты развития и национальные холдинги"
        for c in RECIPIENTS:
            out[f"{sec}.{c}"] = f"{title}: {RECIPIENT_NAMES[c]} (конечный получатель)"
        for f, n in FAMILY_NAMES.items():
            out[f"{sec}.FAM.{f}"] = f"{title}: {n}"
        for h in HOLDINGS:
            out[f"{sec}.VIA.{h}"] = f"{title}: через {RECIPIENT_NAMES[h]} (прямой получатель)"
        for c in REMAINDERS:
            out[f"{sec}.{c}"] = f"{title}: {REMAINDER_NAMES[c]}"
    return out


ITEM_NAMES = item_names()


# ---------------------------------------------------------------- cells and labels

def norm(text) -> str:
    t = " ".join(str(text or "").replace("_x000D_", " ").split()).lower().replace("ё", "е")
    return re.sub(r"^\d{3}\s+", "", t)           # «103 Увеличение …» (2020 editions)


def code(cell) -> int | None:
    if cell in (None, "") or isinstance(cell, bool):
        return None
    try:
        return int(float(str(cell).strip()))
    except ValueError:
        return None


def num(cell) -> float | None:
    if isinstance(cell, (int, float)) and not isinstance(cell, bool):
        return float(cell)
    if isinstance(cell, str) and re.fullmatch(r"\s*-?[\d\s\xa0]+([.,]\d+)?\s*", cell):
        return float(re.sub(r"[\s\xa0]", "", cell).replace(",", "."))
    return None


def classify(label: str, section: str | None = None, abp: int | None = None, prg: int | None = None,
             year: int | None = None) -> dict:
    """{'class': 'DI' | LOCAL_GOV | …, 'direct', 'final', 'family', 'truncated'} of a programme label."""
    n = norm(label)
    hits = sorted((m.start(), m.end(), c) for c, rx in _ENTITIES for m in rx.finditer(n))
    if not hits:
        cls = next((c for c, rx in _OTHER_RULES if rx.search(n)), "OTHER")
        return {"class": cls, "direct": cls, "final": cls, "family": cls, "truncated": False}
    via = VIA_RE.search(n)
    before = [h for h in hits if not via or h[0] < via.start()] or hits
    last = before[-1]
    final = RENAME.get(last[2], last[2])
    truncated = len(" ".join(str(label).split())) >= TRUNCATED_AT and bool(NEXT_RE.search(n[last[1]:]))
    key = (section, abp, prg, year)
    if truncated and key in KNOWN_TRUNCATED:
        final = KNOWN_TRUNCATED[key] or final
        truncated = False
    direct = RENAME.get(hits[0][2], hits[0][2])
    return {"class": "DI", "direct": direct, "final": final, "family": FAMILY[final], "truncated": truncated}


# ---------------------------------------------------------------- periods

_LOOSE_RE = re.compile(r"январ\w*\s*-\s*([а-я]+)\s+(\d{4})", re.IGNORECASE)


def period_of(cell) -> tuple[int, int] | None:
    """minfin._period_of, plus headers without «отчет» («январь-март 2018 г.»: the 2018-03 and
    2019-10 editions)."""
    if not isinstance(cell, str):
        return None
    text = " ".join(cell.replace("_x000D_", " ").split())
    p = minfin._period_of(text)
    if p:
        return p
    m = _LOOSE_RE.search(text)
    if m and minfin._ru_month(m.group(1)):
        return int(m.group(2)), minfin._ru_month(m.group(1))
    return None


def find_period(rows: list[list]) -> tuple[int, int, int] | None:
    for row in rows[:15]:
        for j, cell in enumerate(row or []):
            p = period_of(cell)
            if p:
                return p[0], p[1], j
    return None


def edition_period(sheets: dict[str, list[list]]) -> tuple[int, int] | None:
    """The latest January-to-month header among табл 7, 8 (расх) and 10."""
    cands = [p for name, rows in sheets.items() if minfin._sheet_key(name) in (TABLE7, TABLE8, TABLE10)
             for row in rows[:15] for c in row or [] if (p := period_of(c))]
    return max(cands) if cands else None


# ---------------------------------------------------------------- table 8

def parse_table8(rows: list[list], period: tuple[int, int] | None = None) -> dict:
    """{'period', 'header_period', 'col', 'sections': {LOAN|EQUITY|…: header value}, 'items':
    [programme rows of LOAN and EQUITY]}. `period` overrides the sheet's own (stale) header."""
    per = find_period(rows)
    if per is None:
        raise validation.StructuralChangeError("minfin_quasifiscal: no January-to-month header in «табл 8 (расх)»")
    year, month, col = per
    header_period = (year, month)
    if period is not None:
        year, month = period
    sections: dict[str, float | None] = {}
    items: list[dict] = []
    sec = admin = None
    for i, r in enumerate(rows):
        if not r or len(r) <= col:
            continue
        label = minfin._row_label(r)
        n = norm(label)
        codes = [code(c) for c in (list(r[:4]) + [None] * 4)[:4]]
        v = num(r[col])
        if codes[2] is None and codes[3] is None:
            s = next((k for k, rx in SECTION_RE.items() if rx.match(n)), None)
            if s:
                if s in sections:         # a second block (the January 2014 edition appends Jan-Dec 2013)
                    break
                sec, admin = s, None
                sections[s] = v
                continue
        if sec not in SECTIONS or all(c is None for c in codes):
            continue
        if codes[2] is not None:
            admin = codes[2]
        if codes[3] is None:
            continue
        items.append({"section": sec, "abp": admin, "prg": codes[3], "label": " ".join(label.split()),
                      "value": v or 0.0, "row": i, **classify(label, sec, admin, codes[3], year)})
    return {"period": (year, month), "header_period": header_period, "col": col, "sections": sections, "items": items}


def section_errors(parsed: dict) -> list[str]:
    """Programme rows that do not add up to their section header."""
    out = []
    for sec in SECTIONS:
        leaves = sum(i["value"] for i in parsed["items"] if i["section"] == sec)
        header = parsed["sections"].get(sec)
        if sec not in parsed["sections"] and abs(leaves) < TOLERANCE:     # no such section (the 2013-01 assets)
            continue
        if abs((header or 0.0) - leaves) > TOLERANCE:
            out.append(f"{sec}: programmes sum to {leaves:.3f}, section header {header}")
    return out


def aggregate(parsed: dict) -> dict[str, float]:
    """Every item of ITEM_NAMES for one edition (0 when no programme of it is listed)."""
    out = {k: 0.0 for k in ITEM_NAMES}
    for sec in SECTIONS:
        out[sec] = parsed["sections"].get(sec) or 0.0
    for i in parsed["items"]:
        sec, v = i["section"], i["value"]
        if i["class"] == "DI":
            out[f"{sec}.DI"] += v
            out[f"{sec}.{i['final']}"] += v
            out[f"{sec}.FAM.{i['family']}"] += v
            if i["direct"] in HOLDINGS:
                out[f"{sec}.VIA.{i['direct']}"] += v
        else:
            out[f"{sec}.{i['class']}"] += v
    return {k: round(v, 6) for k, v in out.items()}


def identity_errors(values: dict[str, float]) -> list[str]:
    """SEC = SEC.DI + remainders; SEC.DI = Σ recipients = Σ families."""
    out = []
    for sec in SECTIONS:
        di = values.get(f"{sec}.DI", 0.0)
        checks = {
            f"{sec} = DI + remainders": (values.get(sec, 0.0), di + sum(values.get(f"{sec}.{c}", 0.0) for c in REMAINDERS)),
            f"{sec}.DI = recipients": (di, sum(values.get(f"{sec}.{c}", 0.0) for c in RECIPIENTS)),
            f"{sec}.DI = families": (di, sum(values.get(f"{sec}.FAM.{f}", 0.0) for f in FAMILY_NAMES)),
        }
        out += [f"{k}: {a:.3f} vs {b:.3f}" for k, (a, b) in checks.items() if abs(a - b) > TOLERANCE]
    return out


# ---------------------------------------------------------------- cross-checks (tables 7 and 10)

def table7_totals(rows: list[list], year: int, month: int) -> dict[str, float | None]:
    """«Бюджетные кредиты» / «Приобретение финансовых активов» of табл 7 (republican budget) in
    the column headed with this period; in a December edition the year's report column."""
    col = None
    for r in rows[:10]:
        for j, c in enumerate(r or []):
            if period_of(c) == (year, month):
                col = j
    if col is None and month == 12:
        rx = re.compile(rf"{year}\s*г\.?\s*отчет|отчет\s*{year}")
        col = next((j for r in rows[:10] for j, c in enumerate(r or [])
                    if isinstance(c, str) and rx.search(" ".join(c.split()))), None)
    if col is None:
        return {}
    out: dict[str, float | None] = {}
    for r in rows:
        n = norm(minfin._row_label(r))
        for sec, title in SECTIONS.items():
            if n == title.lower() and sec not in out:
                out[sec] = num(r[col]) if col < len(r) else None
    return out


def table10_di_loans(rows: list[list], year: int, month: int) -> float | None:
    """Economic specifics 513 (specialised organisations) + 519 (other domestic loans) of табл 10,
    or None when the sheet is of another period."""
    per = find_period(rows)
    if per is None or per[:2] != (year, month):
        return None
    col = per[2]
    pats = (r"^бюджетные кредиты специализированным организациям", r"^прочие внутренние бюджетные кредиты")
    found: dict[str, float] = {}
    for r in rows:
        n = norm(minfin._row_label(r))
        for p in pats:
            if p not in found and re.search(p, n):
                found[p] = (num(r[col]) if col < len(r) else None) or 0.0
    return sum(found.values()) if found else None


# ---------------------------------------------------------------- one edition

def edition_records(sheets: dict[str, list[list]]) -> tuple[str | None, list[dict], list[str], list[dict]]:
    """(date, records, warnings, programme rows) of one edition; (None, [], [], []) when it has
    no «табл 8 (расх)». Raises StructuralChangeError when the programmes do not add up."""
    name = next((n for n in sheets if minfin._sheet_key(n) == TABLE8), None)
    if name is None:
        return None, [], [], []
    parsed = parse_table8(sheets[name], period=edition_period(sheets))
    year, month = parsed["period"]
    d = f"{year}-{month:02d}-01"
    errors = section_errors(parsed)
    if errors:
        raise validation.StructuralChangeError(f"minfin_quasifiscal: {d} edition: " + "; ".join(errors))
    values = aggregate(parsed)
    errors = identity_errors(values)
    if errors:
        raise validation.StructuralChangeError(f"minfin_quasifiscal: {d} edition: " + "; ".join(errors))
    warnings = []
    if parsed["header_period"] != parsed["period"]:
        warnings.append(f"{d}: «табл 8 (расх)» headed {parsed['header_period']}, edition {parsed['period']} (stale header)")
    t7 = next((rows for n, rows in sheets.items() if minfin._sheet_key(n) == TABLE7), None)
    for sec, v in (table7_totals(t7, year, month) if t7 else {}).items():
        if v is not None and abs(v - values[sec]) > TOLERANCE:
            known = " (known табл 7 error, source_issues)" if d in KNOWN_TABLE7_ERRORS else ""
            warnings.append(f"{d} {sec}: табл 8 {values[sec]:.3f} vs табл 7 {v:.3f}{known}")
    t10 = next((rows for n, rows in sheets.items() if minfin._sheet_key(n) == TABLE10), None)
    di10 = table10_di_loans(t10, year, month) if t10 else None
    if di10 is not None and abs(di10 - values["LOAN.DI"]) > TOLERANCE:
        warnings.append(f"{d} LOAN.DI: {values['LOAN.DI']:.3f} vs табл 10 513+519 {di10:.3f}")
    for i in parsed["items"]:
        where = f"{d} {i['section']} {i['abp']}/{i['prg']}"
        if i["truncated"]:
            warnings.append(f"{where}: label cut before the final recipient, taken as {i['final']}: …{i['label'][-80:]}")
        if (i["section"] == "LOAN" and i["class"] in ("OTHER_QUASI", "OTHER")
                and not any(re.search(p, norm(i["label"])) for p in KNOWN_OTHER_LOANS)):
            warnings.append(f"{where}: loan to an unknown recipient ({i['class']}): {i['label'][:120]}")
    records = [{"date": d, "region": dims.NATIONAL, "item_code": k, "item_name": ITEM_NAMES[k], "value": v}
               for k, v in values.items()]
    programmes = [{"date": d, **{k: i[k] for k in ("section", "abp", "prg", "class", "direct", "final", "family", "value", "label")}}
                  for i in parsed["items"]]
    return d, records, warnings, programmes


def ytd_warnings(records: list[dict]) -> list[str]:
    """Year-to-date values that fall within a year."""
    by: dict[str, dict[str, float]] = {}
    for r in records:
        by.setdefault(r["item_code"], {})[r["date"]] = float(r["value"])
    out = []
    for item, series in sorted(by.items()):
        prev = None
        for d in sorted(series):
            if prev and prev[0][:4] == d[:4] and series[d] < prev[1] - TOLERANCE:
                out.append(f"{item}: year to date falls from {prev[1]:.1f} ({prev[0][:7]}) to {series[d]:.1f} ({d[:7]})")
            prev = (d, series[d])
    return out


def drop_zero_series(records: list[dict]) -> list[dict]:
    nonzero = {r["item_code"] for r in records if abs(float(r["value"])) > 1e-9}
    return [r for r in records if r["item_code"] in nonzero]


# ---------------------------------------------------------------- fetch

def _documents() -> list[dict]:
    """The bulletins minfin lists, and the untitled editions of UNTITLED_BULLETINS when the
    budget direction's listing still holds them."""
    docs = list(minfin._bulletin_documents())
    have = {str(d.get("id")) for d in docs}
    try:
        listed = minfin._list_documents(directions=minfin.BUDGET_DIRECTION_ID)
    except Exception:  # noqa: BLE001 -- the extra editions are a bonus, not a requirement
        listed = []
    extra = [d for d in listed if str(d.get("id")) in UNTITLED_BULLETINS and str(d.get("id")) not in have
             and d.get("full_text") and str(d["full_text"][0].get("document", "")).lower().endswith((".xlsx", ".xls"))]
    return sorted(docs + extra, key=lambda d: d.get("created_date") or "", reverse=True)


def fetch(ds: dict) -> tuple[list[dict], dict]:
    docs = _documents()
    if not docs:
        raise validation.StructuralChangeError("minfin_quasifiscal: no Statistical Bulletin in the listing")
    got: dict[str, list[dict]] = {}
    warnings: list[str] = []
    for doc in docs:                          # newest first: an edition already read wins
        path = doc["full_text"][0]["document"]
        kind, wb = minfin._open_workbook(minfin._download(path), path)
        names = wb.sheet_names() if kind == "xlrd" else wb.sheetnames
        wanted = [n for n in names if minfin._sheet_key(n) in (TABLE7, TABLE8, TABLE10)]
        if not any(minfin._sheet_key(n) == TABLE8 for n in wanted):
            warnings.append(f"{doc.get('id')}: no «табл 8 (расх)»")
            continue
        sheets = {n: list(minfin._iter_rows(kind, wb, n)) for n in wanted}
        d, records, notes, _ = edition_records(sheets)
        if d and d not in got:
            got[d] = records
            warnings += [f"{doc.get('id')}: {n}" for n in notes]
    if not got:
        raise validation.StructuralChangeError("minfin_quasifiscal: no «табл 8 (расх)» read from any bulletin")
    stored = [{**r, "value": float(r["value"])} for r in dims.load_processed(ds.get("id", DATASET_ID)) if r["date"] not in got]
    records = drop_zero_series([r for recs in got.values() for r in recs] + stored)
    warnings += ytd_warnings(records)
    return sorted(records, key=lambda r: (r["item_code"], r["date"])), {
        "frequency": "monthly", "source_url": minfin.LISTING_URL, "dataset_id": "Statistical Bulletin table 8 (расх)",
        "note": ds.get("note", ""), "warnings": warnings}


# ---------------------------------------------------------------- audit (offline)

def _read_workbook(path: Path) -> dict[str, list[list]]:
    """The tables 7, 8 (расх) and 10 of a local bulletin file (xls or xlsx)."""
    import openpyxl
    import xlrd
    out: dict[str, list[list]] = {}
    if path.suffix.lower() == ".xls" and path.read_bytes()[:4] != b"PK\x03\x04":
        book = xlrd.open_workbook(str(path), on_demand=True)
        for n in book.sheet_names():
            if minfin._sheet_key(n) in (TABLE7, TABLE8, TABLE10):
                sh = book.sheet_by_name(n)
                out[n] = [[sh.cell_value(r, c) for c in range(sh.ncols)] for r in range(sh.nrows)]
    else:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for n in wb.sheetnames:
            if minfin._sheet_key(n) in (TABLE7, TABLE8, TABLE10):
                out[n] = [list(r) for r in wb[n].iter_rows(values_only=True)]
        wb.close()
    return out


def main(argv: list[str] | None = None) -> None:
    """python scripts/fetchers/minfin_quasifiscal.py FILE_OR_DIR … --out programmes.csv

    Reads local bulletin files (e.g. the --cache folder of backfill_minfin_bulletins.py) and
    writes the programme-level audit table (date, section, administrator, programme, class,
    direct, final, family, value, label) plus the checks, without touching data/."""
    import argparse
    import csv
    ap = argparse.ArgumentParser(description=main.__doc__.split("\n\n")[0])
    ap.add_argument("paths", nargs="+", type=Path)
    ap.add_argument("--out", type=Path, required=True)
    args = ap.parse_args(argv)
    files = []
    for p in args.paths:
        files += sorted(f for f in (p.rglob("*") if p.is_dir() else [p]) if f.suffix.lower() in (".xls", ".xlsx"))
    rows, seen = [], set()
    for f in files:
        d, _, notes, programmes = edition_records(_read_workbook(f))
        if d is None or d in seen:
            continue
        seen.add(d)
        rows += [{**p, "file": str(f)} for p in programmes]
        for n in notes:
            print(f"{f.name}: {n}")
    with args.out.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=["date", "section", "abp", "prg", "class", "direct", "final", "family", "value", "label", "file"])
        w.writeheader()
        w.writerows(sorted(rows, key=lambda r: (r["date"], r["section"], r["abp"] or 0, r["prg"])))
    print(f"{len(seen)} editions, {len(rows)} programme rows -> {args.out}")


if __name__ == "__main__":
    main()
