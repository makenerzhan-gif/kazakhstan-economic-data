#!/usr/bin/env python3
"""One-off load of National Fund receipts by tax, 2002-2025, from KGD (2026-09-26).

KGD (the State Revenue Committee of the Ministry of Finance) publishes «Динамика
поступлений налогов и платежей в Национальный фонд»: one workbook per year, rows by
budget classification code, columns January, January-February ... January-December
(thousand KZT, year to date). The page stopped being updated in May 2025 (2025 ends with
January-April) and the file names are irregular (nalogi_rus__16.xlsx, URAS/nalogi_9.xlsx,
statistika/2015/nacfond_2012_po_nalogam.xls), so it is not a daily source: this script
reads the page once, archives every by-tax workbook under data/raw/minfin/ as
NF_RECEIPTS_HISTORY_KGD_<year>, and writes data/reference/nf_receipts_kgd_history.csv.
The daily Minfin fetcher (scripts/fetchers/minfin.py, _fetch_nf) prepends the months
before its first report, and only while KGD equals Minfin on every month both carry.

Rows are read by their Russian label (the code column appears only from 2015). The seven
lines kept are the seven taxes of the Minfin report; corporate income tax is the sum of its
rows (2002-2015 split by taxpayer type, codes 101105/101106/101107). A month is kept only
if the file's own identity holds: all tax rows add up to «ИТОГО по налоговым поступлениям»
(royalty, VAT on production-sharing output and similar pre-2009 lines are in that sum but
not kept). A workbook failing it in any month is left out whole: the 2005 one does (two
corporate income tax rows and the bonus row hold misplaced numbers). There is no 2006
workbook on the page, and 2003-2004 start in August.

Run: python scripts/load_nf_receipts_history.py
"""
from __future__ import annotations

import csv
import html
import io
import re
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl
import requests
import xlrd

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lib import raw_store  # noqa: E402

PAGE_URL = "https://kgd.gov.kz/ru/section/dinamika-postupleniy-nalogov-i-platezhey-v-nacionalnyy-fond"
KGD_BASE = "https://kgd.gov.kz"
HEADERS = {"User-Agent": "Mozilla/5.0 (compatible; KZEconDataPipeline/1.0; +https://github.com/)"}
OUT_PATH = Path(__file__).resolve().parents[1] / "data" / "reference" / "nf_receipts_kgd_history.csv"
TOLERANCE = 2.0  # thousand KZT

RU_MONTHS = {"январ": 1, "феврал": 2, "март": 3, "апрел": 4, "май": 5, "мая": 5, "июн": 6, "июл": 7,
             "август": 8, "сентябр": 9, "октябр": 10, "ноябр": 11, "декабр": 12}
KGD_ROWS = {
    "NF_OIL_CIT_YTD": lambda s: s.startswith("корпор"),
    "NF_EXCESS_PROFIT_TAX_YTD": lambda s: s.startswith("налог на сверхприбыль"),
    "NF_BONUSES_YTD": lambda s: s.startswith("бонус"),
    "NF_MET_YTD": lambda s: s.startswith("налог на добычу полезных"),
    "NF_RENT_TAX_EXPORT_YTD": lambda s: s.startswith("рентный налог на экспорт"),
    "NF_PSA_SHARE_YTD": lambda s: s.startswith("доля рк по разделу"),
    "NF_PSA_ADDITIONAL_PAYMENT_YTD": lambda s: s.startswith(("доп.платеж недропользователя",
                                                              "дополнительный платеж недропользователя")),
}


def _month(text: str) -> int | None:
    t = text.strip().lower()
    return next((n for stem, n in RU_MONTHS.items() if t.startswith(stem)), None)


def list_workbooks(page_html: str) -> dict[int, str]:
    """{year: url} of the «по видам налогов и платежей» workbooks; the year is the last one
    written in the page text just before the link."""
    out: dict[int, str] = {}
    for m in re.finditer(r'<a[^>]+href="([^"]+\.xlsx?)"[^>]*>(.*?)</a>', page_html, re.S | re.I):
        if "по видам" not in html.unescape(m.group(2)):
            continue
        before = html.unescape(re.sub(r"<[^>]+>", " ", page_html[max(0, m.start() - 600):m.start()]))
        years = re.findall(r"\b(20\d\d)\b", before)
        if not years:
            continue
        year = int(years[-1])
        if year in out:
            raise ValueError(f"two by-tax workbooks for {year}: {out[year]} and {m.group(1)}")
        url = m.group(1)
        out[year] = url if url.startswith("http") else KGD_BASE + url
    return out


def workbook_rows(content: bytes) -> list[list]:
    if content.startswith(b"PK"):
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        return [list(r) for r in wb[wb.sheetnames[0]].iter_rows(values_only=True)]
    wb = xlrd.open_workbook(file_contents=content)
    ws = wb.sheet_by_index(0)
    return [[ws.cell_value(r, c) for c in range(ws.ncols)] for r in range(ws.nrows)]


def parse_kgd_sheet(rows: list[list]) -> tuple[dict[str, dict[int, float]], list[int]]:
    """({indicator: {month: thousand KZT}}, months dropped by the identity check)."""
    hi = next(i for i, r in enumerate(rows)                # «январь», «январь-февраль», ...
              if sum(isinstance(c, str) and c.strip().lower().startswith("январ") for c in r) >= 2)
    cols: dict[int, int] = {}
    for j, c in enumerate(rows[hi]):
        if isinstance(c, str) and c.strip().lower().startswith("январ"):
            parts = c.replace(" ", "").split("-")
            cols[j] = _month(parts[-1]) if len(parts) > 1 else 1
    sums: dict[int, float] = {}
    total: dict[int, float] = {}
    got: dict[str, dict[int, float]] = {}                  # only the lines the workbook has
    in_taxes = False
    for r in rows[hi + 1:]:
        label = next((c for c in r if isinstance(c, str) and len(c.strip()) > 3), None)
        if label is None:
            continue
        s = " ".join(label.split()).lower()
        if s.startswith("налоговые поступления"):
            in_taxes = True
            continue
        if s.startswith("итого по налоговым"):
            for j, mo in cols.items():
                v = r[j] if j < len(r) else None
                if isinstance(v, (int, float)):
                    total[mo] = float(v)
            break
        if not in_taxes:
            continue
        for j, mo in cols.items():
            v = r[j] if j < len(r) else None
            if isinstance(v, (int, float)):
                sums[mo] = sums.get(mo, 0.0) + float(v)
        for ind, match in KGD_ROWS.items():
            if match(s):
                line = got.setdefault(ind, {})
                for j, mo in cols.items():
                    v = r[j] if j < len(r) else None
                    if isinstance(v, (int, float)):
                        line[mo] = line.get(mo, 0.0) + float(v)
                break
    bad = sorted(mo for mo in total if abs(sums.get(mo, 0.0) - total[mo]) > TOLERANCE)
    good = set() if bad else set(total)                   # one failing month discards the workbook
    for ind in got:
        # a line the workbook has but leaves blank in a month that passes is zero that month
        got[ind] = {mo: got[ind].get(mo, 0.0) for mo in good}
    return got, bad


def main() -> None:
    today = date.today()
    page = requests.get(PAGE_URL, headers=HEADERS, timeout=60)
    page.raise_for_status()
    books = list_workbooks(page.text)
    out_rows, log = [], {}
    for year, url in sorted(books.items()):
        resp = requests.get(url, headers=HEADERS, timeout=120)
        resp.raise_for_status()
        content = resp.content
        ext = "xlsx" if content.startswith(b"PK") else "xls"
        saved = raw_store.save_raw_bytes("minfin", f"NF_RECEIPTS_HISTORY_KGD_{year}", today, ext, content)
        got, bad = parse_kgd_sheet(workbook_rows(content))
        n = 0
        for ind, months in got.items():
            for mo, v in sorted(months.items()):
                out_rows.append({"indicator_id": ind, "date": f"{year:04d}-{mo:02d}-01", "value_thousand_kzt": round(v, 3),
                                 "source_url": url})
                n += 1
        log[year] = {"url": url, "raw_file": saved.name, "points": n, "months_failing_identity": bad}
        print(year, url, n, "points", f"identity fails in months {bad}" if bad else "")
    out_rows.sort(key=lambda r: (r["indicator_id"], r["date"]))
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with OUT_PATH.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["indicator_id", "date", "value_thousand_kzt", "source_url"], lineterminator="\n")
        w.writeheader()
        w.writerows(out_rows)
    raw_store.write_download_manifest("minfin", "NF_RECEIPTS_HISTORY_KGD", today, {
        "downloaded_at": datetime.now().isoformat(), "source_url": PAGE_URL, "output": str(OUT_PATH.name),
        "workbooks": log})
    print(f"{len(out_rows)} rows -> {OUT_PATH}")


if __name__ == "__main__":
    main()
