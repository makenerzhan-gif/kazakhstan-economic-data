#!/usr/bin/env python3
"""One-off: read every monthly Statistical Bulletin Minfin keeps online (January 2019 on,
91 xlsx on 2026-09-27) into the bulletin series' processed history.

    python scripts/backfill_minfin_bulletins.py [--cache DIR] [--only ID ...] [--yearly 2013-2018]

The daily run sees only the budget direction's latest ~14 bulletins and lays them over the
processed history (fetchers/minfin._with_history); this script supplies that history. It sets
minfin.BULLETIN_ARCHIVE, so the documents come from title searches too, and runs update_minfin
for every series read from the bulletin — the January-to-month ones (_fetch_bulletin_row, state
budget table 3) and the annual ones (_fetch_bulletin_annual_row), then the quarterly state
budget built from them. The archive's workbooks are not copied into data/raw.

--yearly FIRST-LAST: also the years before 2019. Minfin keeps them as one RAR per year («Статистический
бюллетень за 2017 год (12 месяцев)», 12 monthly editions inside; the tables are those of the 2019
editions, same sheet names and labels). Each archive is downloaded into the cache, unpacked with
7z (which must be installed) and its xls/xlsx editions are read like the monthly documents. 2013-2018
are spreadsheets; 2008-2012 are mostly PDF and are not read.

--cache DIR: a folder with the bulletins already downloaded, named <document id>.xlsx/.xls
(downloaded there when missing). Each workbook is read once, in openpyxl's read-only mode, and
its sheets kept in memory for every series (the full-mode load the daily fetcher does would
take over an hour for 91 workbooks x 25 series).
"""
from __future__ import annotations

import argparse
import inspect
import os
import sys
import tempfile
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import update_minfin  # noqa: E402
from fetchers import minfin  # noqa: E402
from lib import pipeline_logging  # noqa: E402

BULLETIN_READERS = ("_fetch_bulletin_row", "_fetch_state_budget_ytd", "_fetch_bulletin_annual_row")


def bulletin_series() -> list[str]:
    ids = [i for i, f in update_minfin.FETCHERS.items()
           if any(r in inspect.getsource(f) for r in BULLETIN_READERS)]
    return [i for i in update_minfin.INDICATOR_IDS if i in ids] + list(minfin.STATE_BUDGET_QUARTERLY)


YEARLY_TITLE = "Статистический бюллетень за {year} год"


def yearly_editions(years: range, cache: Path) -> list[dict]:
    """Pseudo-documents for the xls/xlsx editions inside the yearly RAR archives; the local path
    stands in for the gov.kz file path (install_cache reads it back)."""
    import glob
    import subprocess
    docs = []
    for year in years:
        found = [d for d in minfin._list_documents(title=YEARLY_TITLE.format(year=year), projects="minfin")
                 if d.get("full_text") and str(d["full_text"][0].get("document", "")).lower().endswith(".rar")
                 and str(year) in (d.get("title") or "")]
        if not found:
            print(f"{year}: no yearly archive found")
            continue
        folder = cache / f"yearly_{year}"
        if not folder.exists():
            rar = cache / f"yearly_{year}.rar"
            if not rar.exists():
                rar.write_bytes(minfin._download(found[0]["full_text"][0]["document"]))
            folder.mkdir()
            subprocess.run(["7z", "x", "-y", f"-o{folder}", str(rar)], check=True, capture_output=True)
        files = sorted(f for f in glob.glob(str(folder / "**" / "*"), recursive=True) if f.lower().endswith((".xls", ".xlsx")))
        print(f"{year}: {len(files)} editions from document {found[0]['id']}")
        docs += [{"id": f"yearly:{year}:{Path(f).name}", "title": f"yearly {year} {Path(f).name}", "created_date": f"{year}-12-31",
                  "full_text": [{"document": f}]} for f in files]
    return docs


class _Sheets:
    def __init__(self, sheets: dict[str, list[list]]):
        self.sheets = sheets
        self.sheetnames = list(sheets)

    def sheet_names(self):
        return self.sheetnames


def install_cache(cache: Path) -> None:
    """Downloads go to `cache`; every workbook is parsed once, read-only, for all series."""
    import openpyxl
    import xlrd
    by_path: dict[str, str] = {}
    parsed: dict[str, _Sheets] = {}
    download = minfin._download

    def cached_download(path: str) -> bytes:
        if path.startswith(str(cache)):                   # an edition unpacked from a yearly archive
            return str(Path(path).relative_to(cache)).encode()
        name = by_path.get(path) or path.rstrip("/").split("/")[-1]
        f = cache / name
        if not f.exists():
            f.write_bytes(download(path))
        return f.name.encode()

    def cached_open(content: bytes, hint: str):
        name = content.decode()
        if name not in parsed:
            f = cache / name
            sheets: dict[str, list[list]] = {}
            if f.suffix.lower() == ".xls" and f.read_bytes()[:4] != b"PK\x03\x04":
                book = xlrd.open_workbook(str(f))
                for sh in book.sheets():
                    sheets[sh.name] = [[sh.cell_value(r, c) for c in range(sh.ncols)] for r in range(sh.nrows)]
            else:
                wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
                for ws in wb.worksheets:
                    sheets[ws.title] = [list(r) for r in ws.iter_rows(values_only=True)]
                wb.close()
            parsed[name] = _Sheets(sheets)
        return "sheets", parsed[name]

    def cached_rows(kind, wb, sheet_name=None):
        if kind != "sheets":
            return original_rows(kind, wb, sheet_name)
        return iter(wb.sheets[wb.sheetnames[0] if sheet_name is None else sheet_name])

    for d in minfin._bulletin_documents():
        p = str(d["full_text"][0]["document"])
        by_path[p] = f"{d['id']}{os.path.splitext(p)[1].lower()}"
    original_rows = minfin._iter_rows
    minfin._download = cached_download
    minfin._open_workbook = cached_open
    minfin._iter_rows = cached_rows


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--cache", type=Path, default=Path(tempfile.gettempdir()) / "minfin_bulletins")
    ap.add_argument("--only", nargs="*")
    ap.add_argument("--yearly", help="also the yearly archives, e.g. 2013-2018")
    args = ap.parse_args()
    args.cache.mkdir(parents=True, exist_ok=True)

    minfin.BULLETIN_ARCHIVE = True
    if args.yearly:
        first, last = (int(y) for y in args.yearly.split("-"))
        extra = yearly_editions(range(first, last + 1), args.cache)
        listed = minfin._bulletin_documents
        minfin._bulletin_documents = lambda: listed() + extra      # the monthly documents first (newer editions win)
    docs = minfin._bulletin_documents()
    print(f"{len(docs)} bulletins, {docs[-1].get('created_date', '')[:10]} … {docs[0].get('created_date', '')[:10]}")
    install_cache(args.cache)

    update_minfin.INDICATOR_IDS = args.only or bulletin_series()
    log = pipeline_logging.RunLogger(run_timestamp=datetime.now().isoformat())
    update_minfin.run(log)
    for e in log.entries:
        print(f"{e.dataset:40} {e.status:8} {e.records_processed or 0:4}", (e.errors or [""])[0][:120])


if __name__ == "__main__":
    main()
