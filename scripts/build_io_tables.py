#!/usr/bin/env python3
"""BNS input-output and supply-use tables -> data/reference/io/ (long, gzip CSV).

For the МОБ / CGE (CAEM) work. Not part of the daily run: BNS publishes one edition a
year (the symmetric tables in December, the supply-use tables in November), so run this
after a new edition:

    python scripts/build_io_tables.py

Sources, verified 2026-09-25 (listing pages of stat.gov.kz, rows `divTableRow id=bx_<n>_<element>`):
- «Таблицы "Затраты - Выпуск"» (symmetric, product by product, 68 products), listing
  name=19580, editions for 2021-2024 as xlsx; sheets '1.'-'10.': 1 supply at basic prices,
  2 use at purchasers' prices (with the value-added block), 3 trade margins, 4 transport
  margins, 5 net product taxes, 6 use at basic prices, 7 use of imports, 8 use of domestic
  output, 9 direct requirements (A), 10 total requirements (Leontief inverse). Thousand KZT
  (9 and 10: coefficients). NOTE: BNS's A divides by total RESOURCES at basic prices
  (output + imports), not by output alone -- the column sums of «Итого» + «ВДС» + «Импорт»
  are 1. L = (I - A)^-1 holds to 1e-14 (checked here on every build).
- «Таблицы "Ресурсы - Использование"» (supply-use, 125 products x 72 industries), listing
  name=19765, 2021-2024; sheets '1.'-'8.' as the first eight above. Thousand KZT.
Older years exist only as .rar bulletins (2001-2020) and are not read.

The cover page's «Дата релиза» is stale in some editions (the 2024 file says 22.12.2022);
the listing's publication date is recorded instead.

Output: data/reference/io/{kind}_{year}.csv.gz with columns table, row_no, row_code,
row_name, col_no, col_code, col_name, value -- one row per non-empty numeric cell -- and
data/reference/io/editions.csv listing what was read.
"""
from __future__ import annotations

import csv
import gzip
import html
import io
import re
import sys
from pathlib import Path

import numpy as np
import openpyxl
import requests

REPO_ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = REPO_ROOT / "data" / "reference" / "io"
HEADERS = {"User-Agent": "Mozilla/5.0 (compatible; KZEconDataPipeline/1.0; +https://github.com/)"}
LISTING_URL = "https://stat.gov.kz/ru/industries/economy/national-accounts/spreadsheets/?name={name}"
FILE_URL = "https://stat.gov.kz/api/iblock/element/{eid}/file/ru/"
KINDS = {"io_symmetric": "19580", "supply_use": "19765"}
TITLE_YEAR_RE = re.compile(r"\(\s*(\d{4})\s*г")


def listing(name: str) -> list[dict]:
    page = requests.get(LISTING_URL.format(name=name), headers=HEADERS, timeout=60).text
    out = []
    for chunk in re.split(r'<div class="divTableRow" id="bx_\d+_', page)[1:]:
        eid = chunk.split('"', 1)[0]
        titles = [html.unescape(re.sub(r"\s+", " ", re.sub(r"<[^>]+>", "", t))).strip()
                  for t in re.findall(r'<a href="/api/iblock/element/\d+/file/ru/">(.*?)</a>', chunk, re.S)]
        title = next((t for t in titles if t), "")
        released = re.search(r'text-right">\s*([\d.]+)\s*<', chunk)
        year = TITLE_YEAR_RE.search(title)
        if year:
            out.append({"eid": eid, "title": title, "year": int(year.group(1)),
                        "released": released.group(1) if released else ""})
    return out


def parse_sheet(rows: list[list]) -> list[dict]:
    """Cells of one table sheet: the header row is the one whose first cell is '№'; column
    names sit two rows above it and product codes one row above; data rows follow, with
    the row number, code and name in the first three columns."""
    hi = next((i for i, r in enumerate(rows) if r and str(r[0]).strip() == "№"), None)
    if hi is None or hi < 2:
        return []
    names, codes, numbers = rows[hi - 2], rows[hi - 1], rows[hi]
    out = []
    for r in rows[hi + 1:]:
        if not r or r[0] in (None, "") or r[2] in (None, ""):
            continue
        for j in range(3, len(r)):
            v = r[j]
            if isinstance(v, str):
                try:
                    v = float(v.replace(" ", "").replace(",", "."))
                except ValueError:
                    continue
            if v is None or not isinstance(v, (int, float)):
                continue
            out.append({"row_no": str(r[0]).strip(), "row_code": str(r[1] or "").strip(),
                        "row_name": " ".join(str(r[2]).split()),
                        "col_no": str(numbers[j] if j < len(numbers) and numbers[j] is not None else j).strip(),
                        "col_code": str(codes[j] or "").strip() if j < len(codes) else "",
                        "col_name": " ".join(str(names[j] or "").split()) if j < len(names) else "",
                        "value": float(v)})
    return out


def check_leontief(cells: list[dict]) -> float:
    """max |L - (I - A)^-1| over the product block of tables 9 and 10."""
    def matrix(table: str):
        block = [c for c in cells if c["table"] == table and c["row_code"] and c["col_code"]]
        codes = [c["row_code"] for c in block if c["col_code"] == block[0]["col_code"]]
        index = {k: i for i, k in enumerate(codes)}
        m = np.zeros((len(codes), len(codes)))
        for c in block:
            if c["row_code"] in index and c["col_code"] in index:
                m[index[c["row_code"]], index[c["col_code"]]] = c["value"]
        return m
    a, l = matrix("9"), matrix("10")
    return float(np.abs(np.linalg.inv(np.eye(len(a)) - a) - l).max())


def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    editions = []
    for kind, name in KINDS.items():
        for ed in listing(name):
            if not ed["title"] or ed["year"] < 2021:
                continue
            content = requests.get(FILE_URL.format(eid=ed["eid"]), headers=HEADERS, timeout=120).content
            if content[:2] != b"PK":
                print(f"{kind} {ed['year']}: element {ed['eid']} is not an xlsx, skipped", file=sys.stderr)
                continue
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            cells = []
            for sheet in wb.sheetnames:
                m = re.fullmatch(r"\s*(\d+)\.\s*", sheet)
                if not m:
                    continue
                for c in parse_sheet([list(r) for r in wb[sheet].iter_rows(values_only=True)]):
                    cells.append({"table": m.group(1), **c})
            check = ""
            if kind == "io_symmetric":
                err = check_leontief(cells)
                if err > 1e-8:
                    raise SystemExit(f"{kind} {ed['year']}: L differs from (I-A)^-1 by {err:g}")
                check = f"max|L-(I-A)^-1|={err:.1e}"
            path = OUT_DIR / f"{kind}_{ed['year']}.csv.gz"
            with gzip.open(path, "wt", newline="", encoding="utf-8") as f:
                w = csv.DictWriter(f, fieldnames=["table", "row_no", "row_code", "row_name", "col_no", "col_code",
                                                  "col_name", "value"])
                w.writeheader()
                w.writerows(cells)
            editions.append({"kind": kind, "year": ed["year"], "element_id": ed["eid"], "released": ed["released"],
                             "title": ed["title"], "cells": len(cells), "check": check})
            print(f"{kind} {ed['year']}: {len(cells)} cells, element {ed['eid']} {check}")
    with (OUT_DIR / "editions.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["kind", "year", "element_id", "released", "title", "cells", "check"])
        w.writeheader()
        w.writerows(sorted(editions, key=lambda e: (e["kind"], e["year"])))
    return 0


if __name__ == "__main__":
    sys.exit(main())
